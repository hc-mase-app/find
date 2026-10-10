from __future__ import annotations

import base64
import csv
import io
import json
import os
import re
import sqlite3
import sys
import tempfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


APP_DIR = Path(sys.executable).resolve().parent if getattr(sys, "frozen", False) else Path(__file__).resolve().parent
DATA_DIR = APP_DIR / "data"
DB_PATH = DATA_DIR / "imei_pos.db"
BACKUP_DIR = APP_DIR / "backup"
DATABASE_URL = os.environ.get("DATABASE_URL", "").strip()
USE_POSTGRES = DATABASE_URL.startswith(("postgres://", "postgresql://"))
HOST = "127.0.0.1"
PORT = 8765


class DatabaseConnection:
    """Small compatibility layer so the same application works with SQLite or Neon Postgres."""

    def __init__(self) -> None:
        self.backend = "postgres" if USE_POSTGRES else "sqlite"
        if USE_POSTGRES:
            import psycopg
            from psycopg.rows import dict_row

            self.raw = psycopg.connect(DATABASE_URL, row_factory=dict_row)
        else:
            DATA_DIR.mkdir(parents=True, exist_ok=True)
            self.raw = sqlite3.connect(DB_PATH)
            self.raw.row_factory = sqlite3.Row
            self.raw.execute("PRAGMA foreign_keys = ON")

    def _sql(self, statement: str) -> str:
        if self.backend == "sqlite":
            return statement
        converted = statement.replace("?", "%s")
        converted = converted.replace("date(%s)", "CAST(%s AS DATE)")
        converted = converted.replace(
            "CAST(julianday('now') - julianday(substr(b.tanggal_masuk, 1, 10)) AS INTEGER)",
            "CAST(CURRENT_DATE - CAST(SUBSTRING(b.tanggal_masuk, 1, 10) AS DATE) AS INTEGER)",
        )
        return converted

    def execute(self, statement: str, params: Any = ()) -> Any:
        return self.raw.execute(self._sql(statement), params)

    def executescript(self, script: str) -> None:
        if self.backend == "sqlite":
            self.raw.executescript(script)
            return
        for statement in script.split(";"):
            if statement.strip():
                self.raw.execute(statement)

    def table_columns(self, table_name: str) -> set[str]:
        if self.backend == "sqlite":
            return {row["name"] for row in self.execute(f"PRAGMA table_info({table_name})").fetchall()}
        rows = self.execute(
            "SELECT column_name AS name FROM information_schema.columns WHERE table_schema='public' AND table_name=?",
            (table_name,),
        ).fetchall()
        return {row["name"] for row in rows}

    def insert_and_get_id(self, statement: str, params: Any) -> int:
        if self.backend == "postgres":
            row = self.execute(statement.rstrip().rstrip(";") + " RETURNING id", params).fetchone()
            return int(row["id"])
        return int(self.execute(statement, params).lastrowid)

    def __enter__(self) -> "DatabaseConnection":
        return self

    def __exit__(self, exc_type: Any, exc: Any, traceback: Any) -> None:
        try:
            if exc_type is None:
                self.raw.commit()
            else:
                self.raw.rollback()
        finally:
            self.raw.close()


@dataclass(frozen=True)
class BackupReference:
    name: str


def now_text() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def clean_imei(value: Any) -> str:
    text = clean_text(value)
    text = re.sub(r"\s+", "", text)
    if text.endswith(".0"):
        text = text[:-2]
    return text


def to_number(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    text = re.sub(r"[^\d,.\-]", "", text)
    if not text:
        return None
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def rupiah(value: Any) -> str:
    number = to_number(value)
    if number is None:
        return "-"
    return f"Rp {number:,.0f}".replace(",", ".")


def db() -> DatabaseConnection:
    return DatabaseConnection()


def init_db() -> None:
    with db() as con:
        id_column = "BIGSERIAL PRIMARY KEY" if con.backend == "postgres" else "INTEGER PRIMARY KEY AUTOINCREMENT"
        con.executescript(
            f"""
            CREATE TABLE IF NOT EXISTS barang (
                imei TEXT PRIMARY KEY,
                tipe TEXT NOT NULL,
                hpp REAL,
                modal_awal REAL,
                grade TEXT,
                spek TEXT,
                keterangan TEXT,
                status TEXT NOT NULL DEFAULT 'BELUM_LAKU',
                tanggal_masuk TEXT,
                sumber_file TEXT,
                sumber_sheet TEXT,
                updated_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS penjualan (
                id {id_column},
                imei TEXT NOT NULL,
                tipe TEXT NOT NULL,
                hpp REAL,
                harga_jual REAL NOT NULL,
                laba_rugi REAL,
                customer TEXT,
                catatan TEXT,
                batch_id TEXT,
                tanggal_jual TEXT NOT NULL,
                created_at TEXT NOT NULL,
                FOREIGN KEY (imei) REFERENCES barang(imei)
            );

            CREATE TABLE IF NOT EXISTS import_log (
                id {id_column},
                nama_file TEXT NOT NULL,
                import_type TEXT DEFAULT 'DATA_MASUK',
                supplier TEXT,
                invoice TEXT,
                tanggal_pembelian TEXT,
                jumlah_baris INTEGER NOT NULL,
                barang_baru INTEGER NOT NULL,
                barang_update INTEGER NOT NULL,
                dilewati INTEGER NOT NULL,
                imported_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS audit_log (
                id {id_column},
                aksi TEXT NOT NULL,
                imei TEXT,
                detail TEXT,
                created_at TEXT NOT NULL
            );

            CREATE TABLE IF NOT EXISTS purchase_batches (
                id {id_column},
                batch_id TEXT NOT NULL UNIQUE,
                source_name TEXT,
                supplier TEXT,
                invoice TEXT,
                tanggal_pembelian TEXT,
                jumlah_baris INTEGER NOT NULL DEFAULT 0,
                total_modal REAL NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            );

            CREATE INDEX IF NOT EXISTS idx_barang_status ON barang(status);
            CREATE INDEX IF NOT EXISTS idx_barang_tipe ON barang(tipe);
            CREATE INDEX IF NOT EXISTS idx_penjualan_tanggal ON penjualan(tanggal_jual);
            """
        )
        if con.backend == "postgres":
            con.execute(
                f"""
                CREATE TABLE IF NOT EXISTS database_backups (
                    id {id_column},
                    label TEXT NOT NULL,
                    payload TEXT NOT NULL,
                    created_at TEXT NOT NULL
                )
                """
            )
        existing_columns = con.table_columns("import_log")
        for column_name, column_sql in {
            "import_type": "TEXT DEFAULT 'DATA_MASUK'",
            "supplier": "TEXT",
            "invoice": "TEXT",
            "tanggal_pembelian": "TEXT",
        }.items():
            if column_name not in existing_columns:
                con.execute(f"ALTER TABLE import_log ADD COLUMN {column_name} {column_sql}")
        penjualan_columns = con.table_columns("penjualan")
        if "batch_id" not in penjualan_columns:
            con.execute("ALTER TABLE penjualan ADD COLUMN batch_id TEXT")
        barang_columns = con.table_columns("barang")
        for column_name, column_sql in {
            "purchase_batch_id": "TEXT",
            "supplier": "TEXT",
            "invoice": "TEXT",
        }.items():
            if column_name not in barang_columns:
                con.execute(f"ALTER TABLE barang ADD COLUMN {column_name} {column_sql}")
        con.execute("CREATE INDEX IF NOT EXISTS idx_barang_purchase_batch ON barang(purchase_batch_id)")


def row_values(ws, row_index: int) -> list[Any]:
    return [cell.value for cell in ws[row_index]]


def find_header(ws) -> tuple[int | None, dict[str, int]]:
    keywords = {
        "imei": ("IMEI", "SERIAL", "SN"),
        "tipe": ("TIPE", "TYPE", "BARANG", "NAMA"),
        "grade": ("GRADE",),
        "spek": ("SPEK", "SPEC"),
        "hpp": ("MODAL POT", "HPP", "HARGA POKOK", "POKOK"),
        "modal_awal": ("MODAL", "HARGA SEBELUM DISKON", "HARGA"),
        "keterangan": ("KETERANGAN", "KET", "NOTE", "CATATAN"),
    }
    for ridx in range(1, min(ws.max_row, 25) + 1):
        values = [clean_text(v).upper() for v in row_values(ws, ridx)]
        joined = " ".join(values)
        if "IMEI" not in joined or ("TYPE" not in joined and "TIPE" not in joined):
            continue
        mapping: dict[str, int] = {}
        for cidx, value in enumerate(values):
            for name, keys in keywords.items():
                if name in mapping:
                    continue
                if name == "modal_awal" and "MODAL POT" in value:
                    continue
                if any(key in value for key in keys):
                    mapping[name] = cidx
        # "Modal Pot" adalah sumber HPP utama, meskipun file juga memiliki kolom HPP lain.
        for cidx, value in enumerate(values):
            if value == "MODAL POT" or value.startswith("MODAL POT "):
                mapping["hpp"] = cidx
                break
        if "imei" in mapping and "tipe" in mapping:
            return ridx, mapping
    return None, {}


def infer_record_from_no_header(row: tuple[Any, ...]) -> dict[str, Any] | None:
    values = list(row)
    if len(values) < 2:
        return None

    # Pattern: IMEI, Tipe, Grade, Spek, Modal awal, HPP/modal pot, Harga jual, Laba
    if values[0] is not None and values[1] is not None and re.search(r"[A-Za-z]", clean_text(values[1])):
        return {
            "imei": clean_imei(values[0]),
            "tipe": clean_text(values[1]),
            "grade": clean_text(values[2]) if len(values) > 2 else "",
            "spek": clean_text(values[3]) if len(values) > 3 else "",
            "modal_awal": to_number(values[4]) if len(values) > 4 else None,
            "hpp": to_number(values[5]) if len(values) > 5 else to_number(values[4]) if len(values) > 4 else None,
            "keterangan": clean_text(values[8]) if len(values) > 8 else "",
        }

    # Pattern: NO, IMEI, Tipe, Grade, Spek, Modal awal, HPP/modal pot, Keterangan
    if len(values) >= 3 and values[1] is not None and values[2] is not None and re.search(r"[A-Za-z]", clean_text(values[2])):
        return {
            "imei": clean_imei(values[1]),
            "tipe": clean_text(values[2]),
            "grade": clean_text(values[3]) if len(values) > 3 else "",
            "spek": clean_text(values[4]) if len(values) > 4 else "",
            "modal_awal": to_number(values[5]) if len(values) > 5 else None,
            "hpp": to_number(values[6]) if len(values) > 6 else to_number(values[5]) if len(values) > 5 else None,
            "keterangan": clean_text(values[7]) if len(values) > 7 else "",
        }

    return None


def parse_excel(path: Path, original_name: str) -> list[dict[str, Any]]:
    wb = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []

    for ws in wb.worksheets:
        if ws.title.strip().lower() in {"petunjuk", "readme", "instruksi", "panduan"}:
            continue
        header_row, mapping = find_header(ws)
        start_row = (header_row + 1) if header_row else 1

        for row_number, row in enumerate(ws.iter_rows(min_row=start_row, values_only=True), start=start_row):
            if not any(value is not None for value in row):
                continue

            if mapping:
                def get(name: str) -> Any:
                    idx = mapping.get(name)
                    return row[idx] if idx is not None and idx < len(row) else None

                hpp_value = to_number(get("hpp"))
                if "hpp" in mapping and hpp_value is None:
                    raise ValueError(
                        f"Nilai Modal Pot/HPP kosong pada sheet {ws.title}, baris {row_number}, "
                        f"IMEI {clean_imei(get('imei')) or '-'}. Lengkapi nilainya sebelum upload."
                    )
                record = {
                    "imei": clean_imei(get("imei")),
                    "tipe": clean_text(get("tipe")),
                    "grade": clean_text(get("grade")),
                    "spek": clean_text(get("spek")),
                    "modal_awal": to_number(get("modal_awal")),
                    "hpp": hpp_value if "hpp" in mapping else to_number(get("modal_awal")),
                    "keterangan": clean_text(get("keterangan")),
                }
            else:
                record = infer_record_from_no_header(row)

            if not record:
                continue

            imei = record["imei"]
            tipe = record["tipe"]
            if not imei or not tipe:
                continue
            if imei.upper() in {"IMEI", "NO"} or tipe.upper() in {"TYPE", "TIPE"}:
                continue

            record["sumber_file"] = original_name
            record["sumber_sheet"] = ws.title
            records.append(record)

    # Deduplicate in one import by IMEI, keeping the last row seen.
    deduped: dict[str, dict[str, Any]] = {}
    for record in records:
        deduped[record["imei"]] = record
    return list(deduped.values())


def parse_csv_file(path: Path, original_name: str) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8-sig")
    sample = text[:2048]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    records: list[dict[str, Any]] = []
    for row_number, row in enumerate(reader, start=2):
        upper_map = {clean_text(k).upper(): v for k, v in row.items() if k is not None}
        imei = upper_map.get("IMEI") or upper_map.get("SERIAL") or upper_map.get("SN")
        tipe = upper_map.get("TIPE") or upper_map.get("TYPE") or upper_map.get("BARANG") or upper_map.get("NAMA")
        if "MODAL POT" in upper_map:
            hpp = upper_map.get("MODAL POT")
        else:
            hpp = upper_map.get("HPP") or upper_map.get("HARGA POKOK") or upper_map.get("MODAL")
        record = {
            "imei": clean_imei(imei),
            "tipe": clean_text(tipe),
            "hpp": to_number(hpp),
            "modal_awal": to_number(upper_map.get("MODAL")) or to_number(hpp),
            "grade": clean_text(upper_map.get("GRADE")),
            "spek": clean_text(upper_map.get("SPEK") or upper_map.get("SPEC")),
            "keterangan": clean_text(upper_map.get("KETERANGAN") or upper_map.get("KET") or upper_map.get("CATATAN")),
            "sumber_file": original_name,
            "sumber_sheet": "CSV",
        }
        if record["imei"] and record["tipe"] and record["hpp"] is None:
            raise ValueError(
                f"Nilai Modal Pot/HPP kosong pada CSV baris {row_number}, IMEI {record['imei']}. "
                "Lengkapi nilainya sebelum upload."
            )
        if record["imei"] and record["tipe"]:
            records.append(record)
    deduped: dict[str, dict[str, Any]] = {}
    for record in records:
        deduped[record["imei"]] = record
    return list(deduped.values())


def parse_upload_file(path: Path, original_name: str) -> list[dict[str, Any]]:
    suffix = Path(original_name).suffix.lower()
    if suffix == ".csv":
        return parse_csv_file(path, original_name)
    return parse_excel(path, original_name)


def import_records(records: list[dict[str, Any]], filename: str, metadata: dict[str, Any] | None = None) -> dict[str, Any]:
    metadata = metadata or {}
    baru = update = dilewati = 0
    timestamp = now_text()
    batch_id = f"BM-{datetime.now().strftime('%Y%m%d%H%M%S%f')}"
    supplier = clean_text(metadata.get("supplier"))
    invoice = clean_text(metadata.get("invoice"))
    tanggal_pembelian = clean_text(metadata.get("tanggal_pembelian")) or timestamp[:10]
    total_modal = sum((to_number(record.get("hpp")) or 0) for record in records)
    with db() as con:
        con.execute(
            """
            INSERT INTO purchase_batches
            (batch_id, source_name, supplier, invoice, tanggal_pembelian, jumlah_baris, total_modal, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (batch_id, filename, supplier, invoice, tanggal_pembelian, len(records), total_modal, timestamp),
        )
        for record in records:
            existing = con.execute("SELECT status FROM barang WHERE imei = ?", (record["imei"],)).fetchone()
            if existing:
                update += 1
                con.execute(
                    """
                    UPDATE barang
                    SET tipe = ?, hpp = ?, modal_awal = ?, grade = ?, spek = ?, keterangan = ?,
                        tanggal_masuk = COALESCE(tanggal_masuk, ?),
                        sumber_file = ?, sumber_sheet = ?, updated_at = ?, purchase_batch_id = ?, supplier = ?, invoice = ?
                    WHERE imei = ?
                    """,
                    (
                        record["tipe"],
                        record["hpp"],
                        record["modal_awal"],
                        record["grade"],
                        record["spek"],
                        record["keterangan"],
                        timestamp,
                        record["sumber_file"],
                        record["sumber_sheet"],
                        timestamp,
                        batch_id,
                        supplier,
                        invoice,
                        record["imei"],
                    ),
                )
            else:
                baru += 1
                con.execute(
                    """
                    INSERT INTO barang
                    (imei, tipe, hpp, modal_awal, grade, spek, keterangan, status, tanggal_masuk,
                     sumber_file, sumber_sheet, updated_at, purchase_batch_id, supplier, invoice)
                    VALUES (?, ?, ?, ?, ?, ?, ?, 'BELUM_LAKU', ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        record["imei"],
                        record["tipe"],
                        record["hpp"],
                        record["modal_awal"],
                        record["grade"],
                        record["spek"],
                        record["keterangan"],
                        timestamp,
                        record["sumber_file"],
                        record["sumber_sheet"],
                        timestamp,
                        batch_id,
                        supplier,
                        invoice,
                    ),
                )

        con.execute(
            """
            INSERT INTO import_log
            (nama_file, import_type, supplier, invoice, tanggal_pembelian, jumlah_baris,
             barang_baru, barang_update, dilewati, imported_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                filename,
                clean_text(metadata.get("import_type")) or "DATA_MASUK",
                supplier,
                invoice,
                tanggal_pembelian,
                len(records),
                baru,
                update,
                dilewati,
                timestamp,
            ),
        )
    return {"jumlah_baris": len(records), "barang_baru": baru, "barang_update": update, "dilewati": dilewati, "batch_id": batch_id}


def row_to_dict(row: Any | None) -> dict[str, Any] | None:
    if not row:
        return None
    return {key: row[key] for key in row.keys()}


def find_items_by_imei_input(imei_input: str, only_unsold: bool = False, limit: int = 20) -> list[dict[str, Any]]:
    query = clean_imei(imei_input)
    if not query:
        return []

    with db() as con:
        params: list[Any]
        if len(query) >= 10:
            sql = "SELECT * FROM barang WHERE imei = ?"
            params = [query]
            if only_unsold:
                sql += " AND status = 'BELUM_LAKU'"
            sql += " ORDER BY status, tipe, imei LIMIT ?"
            params.append(limit)
        else:
            sql = "SELECT * FROM barang WHERE imei LIKE ?"
            params = [f"%{query}"]
            if only_unsold:
                sql += " AND status = 'BELUM_LAKU'"
            sql += " ORDER BY status, tipe, imei LIMIT ?"
            params.append(limit)
        rows = con.execute(sql, params).fetchall()
    return [row_to_dict(row) for row in rows if row]


def resolve_single_item_for_sale(imei_input: str) -> dict[str, Any]:
    matches = find_items_by_imei_input(imei_input, only_unsold=True, limit=10)
    if not matches:
        raise ValueError("IMEI tidak ditemukan atau barang sudah terjual.")
    if len(matches) > 1:
        choices = ", ".join(f"{item['imei']} ({item['tipe']})" for item in matches[:5])
        raise ValueError(f"6 digit ini cocok ke lebih dari 1 barang. Pilih IMEI lengkap: {choices}")
    return matches[0]


def summary(start_date: str = "", end_date: str = "") -> dict[str, Any]:
    with db() as con:
        stok = con.execute("SELECT COUNT(*) c FROM barang WHERE status = 'BELUM_LAKU'").fetchone()["c"]
        terjual = con.execute("SELECT COUNT(*) c FROM barang WHERE status = 'TERJUAL'").fetchone()["c"]
        filters = ["1=1"]
        params: list[Any] = []
        if start_date:
            filters.append("date(tanggal_jual) >= date(?)")
            params.append(start_date)
        if end_date:
            filters.append("date(tanggal_jual) <= date(?)")
            params.append(end_date)
        where = " WHERE " + " AND ".join(filters)
        totals = con.execute(
            f"SELECT COUNT(*) jumlah, COALESCE(SUM(hpp), 0) total_hpp, COALESCE(SUM(harga_jual), 0) total_jual, COALESCE(SUM(laba_rugi), 0) total_laba FROM penjualan{where}",
            params,
        ).fetchone()
        type_rows = con.execute(
            f"SELECT tipe, COUNT(*) jumlah FROM penjualan{where} GROUP BY tipe",
            params,
        ).fetchall()
        nilai_stok = con.execute("SELECT COALESCE(SUM(hpp), 0) s FROM barang WHERE status='BELUM_LAKU'").fetchone()["s"]

    brand_counts: dict[str, int] = {}
    top_types: list[dict[str, Any]] = []
    for row in type_rows:
        brand = brand_from_tipe(row["tipe"])
        count = int(row["jumlah"])
        brand_counts[brand] = brand_counts.get(brand, 0) + count
        top_types.append({"tipe": row["tipe"], "brand": brand, "count": count})

    top_brands = [
        {"brand": brand, "count": count}
        for brand, count in sorted(brand_counts.items(), key=lambda item: (-item[1], item[0]))[:5]
    ]
    top_types.sort(key=lambda item: (-item["count"], item["tipe"]))
    return {
        "stok": stok,
        "terjual": terjual,
        "jumlah_terjual_periode": totals["jumlah"],
        "total_hpp": totals["total_hpp"],
        "total_laba": totals["total_laba"],
        "total_jual": totals["total_jual"],
        "nilai_stok": nilai_stok,
        "top_brands": top_brands,
        "top_types": top_types[:5],
    }


def list_barang(status: str | None = None, limit: int = 200) -> list[dict[str, Any]]:
    with db() as con:
        if status:
            rows = con.execute(
                "SELECT * FROM barang WHERE status = ? ORDER BY updated_at DESC LIMIT ?",
                (status, limit),
            ).fetchall()
        else:
            rows = con.execute("SELECT * FROM barang ORDER BY updated_at DESC LIMIT ?", (limit,)).fetchall()
    return [row_to_dict(row) for row in rows if row]


def list_penjualan(limit: int = 200, start_date: str = "", end_date: str = "", query: str = "") -> list[dict[str, Any]]:
    with db() as con:
        filters = ["1=1"]
        params: list[Any] = []
        if start_date:
            filters.append("date(tanggal_jual) >= date(?)")
            params.append(start_date)
        if end_date:
            filters.append("date(tanggal_jual) <= date(?)")
            params.append(end_date)
        if query:
            filters.append("(imei LIKE ? OR tipe LIKE ? OR customer LIKE ? OR batch_id LIKE ? OR CAST(id AS TEXT) LIKE ?)")
            needle = f"%{query}%"
            params.extend([needle] * 5)
        params.append(limit)
        rows = con.execute(
            f"SELECT * FROM penjualan WHERE {' AND '.join(filters)} ORDER BY tanggal_jual DESC, id DESC LIMIT ?",
            params,
        ).fetchall()
    items = [row_to_dict(row) for row in rows if row]
    for item in items:
        item["brand"] = brand_from_tipe(item.get("tipe"))
    return items


def list_purchase_batches(limit: int = 200) -> list[dict[str, Any]]:
    with db() as con:
        rows = con.execute(
            "SELECT * FROM purchase_batches ORDER BY created_at DESC, id DESC LIMIT ?", (limit,)
        ).fetchall()
    return [row_to_dict(row) for row in rows if row]


def list_inventory(status: str = "", query: str = "", limit: int = 1000) -> list[dict[str, Any]]:
    filters = ["1=1"]
    params: list[Any] = []
    if status:
        filters.append("b.status = ?")
        params.append(status)
    if query:
        filters.append("(b.imei LIKE ? OR b.tipe LIKE ? OR b.supplier LIKE ?)")
        needle = f"%{query}%"
        params.extend([needle, needle, needle])
    params.append(limit)
    with db() as con:
        rows = con.execute(
            f"""
            SELECT b.*, CAST(julianday('now') - julianday(substr(b.tanggal_masuk, 1, 10)) AS INTEGER) AS umur_hari
            FROM barang b
            WHERE {' AND '.join(filters)}
            ORDER BY b.status, b.updated_at DESC
            LIMIT ?
            """,
            params,
        ).fetchall()
    items = [row_to_dict(row) for row in rows if row]
    for item in items:
        item["brand"] = brand_from_tipe(item.get("tipe"))
    return items


def inventory_data(status: str = "", query: str = "", brand: str = "", limit: int = 2000) -> dict[str, Any]:
    items = list_inventory(status, query, 10000)
    brand_counts: dict[str, int] = {}
    for item in items:
        item_brand = item.get("brand") or "Lainnya"
        brand_counts[item_brand] = brand_counts.get(item_brand, 0) + 1

    brands = [
        {"brand": item_brand, "count": count}
        for item_brand, count in sorted(brand_counts.items(), key=lambda item: (-item[1], item[0]))
    ]
    valid_brands = set(brand_counts)
    selected_brand = brand if brand in valid_brands else ""
    if selected_brand:
        items = [item for item in items if item.get("brand") == selected_brand]
    return {"items": items[:limit], "brands": brands, "selected_brand": selected_brand}


def imei_history(imei_input: str) -> dict[str, Any]:
    matches = find_items_by_imei_input(imei_input, only_unsold=False, limit=10)
    if not matches:
        raise ValueError("IMEI tidak ditemukan.")
    if len(matches) > 1:
        raise ValueError("Input cocok ke beberapa IMEI. Gunakan IMEI lengkap.")
    item = matches[0]
    imei = item["imei"]
    events: list[dict[str, Any]] = []
    if item.get("tanggal_masuk"):
        events.append({"tanggal": item["tanggal_masuk"], "jenis": "BARANG_MASUK", "detail": f"{item.get('supplier') or '-'} / {item.get('invoice') or '-'}"})
    with db() as con:
        for row in con.execute("SELECT * FROM penjualan WHERE imei=? ORDER BY created_at", (imei,)):
            events.append({"tanggal": row["created_at"], "jenis": "PENJUALAN", "detail": f"Harga {rupiah(row['harga_jual'])}; customer {row['customer'] or '-'}"})
        for row in con.execute("SELECT * FROM audit_log WHERE imei=? ORDER BY created_at", (imei,)):
            events.append({"tanggal": row["created_at"], "jenis": row["aksi"], "detail": row["detail"]})
    events.sort(key=lambda event: event.get("tanggal") or "")
    return {"item": item, "events": events}


def brand_from_tipe(tipe: str | None) -> str:
    text = clean_text(tipe).upper()
    if not text:
        return "Lainnya"
    rules = [
        ("IPHONE", "iPhone"),
        ("IPAD", "iPad"),
        ("SAMSUNG", "Samsung"),
        ("OPPO", "Oppo"),
        ("VIVO", "Vivo"),
        ("XIAOMI", "Xiaomi"),
        ("REDMI", "Xiaomi"),
        ("POCO", "Xiaomi"),
        ("INFINIX", "Infinix"),
        ("REALME", "Realme"),
        ("HUAWEI", "Huawei"),
        ("GARMIN", "Garmin"),
        ("ITEL", "Itel"),
        ("NOKIA", "Nokia"),
    ]
    for needle, brand in rules:
        if needle in text:
            return brand
    return clean_text(tipe).split()[0].title() if clean_text(tipe) else "Lainnya"


def group_summary(status: str = "BELUM_LAKU") -> list[dict[str, Any]]:
    with db() as con:
        rows = con.execute("SELECT tipe, COUNT(*) count FROM barang WHERE status = ? GROUP BY tipe", (status,)).fetchall()
    grouped: dict[str, int] = {}
    for row in rows:
        brand = brand_from_tipe(row["tipe"])
        grouped[brand] = grouped.get(brand, 0) + int(row["count"])
    return [
        {"brand": brand, "count": count}
        for brand, count in sorted(grouped.items(), key=lambda item: (-item[1], item[0]))
    ]


def type_summary(status: str = "BELUM_LAKU", brand: str = "iPhone") -> list[dict[str, Any]]:
    with db() as con:
        rows = con.execute("SELECT tipe, COUNT(*) count FROM barang WHERE status = ? GROUP BY tipe", (status,)).fetchall()
    result = [
        {"tipe": row["tipe"], "count": int(row["count"])}
        for row in rows
        if brand_from_tipe(row["tipe"]) == brand
    ]
    return sorted(result, key=lambda item: (-item["count"], item["tipe"]))[:50]


def type_items(status: str = "BELUM_LAKU", tipe: str = "", limit: int = 100) -> list[dict[str, Any]]:
    with db() as con:
        rows = con.execute(
            """
            SELECT b.*,
                   p.id AS sale_id,
                   p.harga_jual,
                   p.laba_rugi,
                   p.customer,
                   p.batch_id,
                   p.tanggal_jual
            FROM barang b
            LEFT JOIN penjualan p ON p.id = (
                SELECT p2.id
                FROM penjualan p2
                WHERE p2.imei = b.imei
                ORDER BY p2.id DESC
                LIMIT 1
            )
            WHERE b.status = ? AND b.tipe = ?
            ORDER BY b.grade, b.imei
            LIMIT ?
            """,
            (status, tipe, limit),
        ).fetchall()
    return [row_to_dict(row) for row in rows if row]


def create_database_backup(label: str) -> Path | BackupReference:
    if USE_POSTGRES:
        snapshot: dict[str, list[dict[str, Any]]] = {}
        with db() as con:
            for table_name in ("barang", "penjualan", "import_log", "audit_log", "purchase_batches"):
                snapshot[table_name] = [row_to_dict(row) or {} for row in con.execute(f"SELECT * FROM {table_name}").fetchall()]
            backup_name = f"neon-{datetime.now().strftime('%Y%m%d-%H%M%S-%f')}-{label}"
            con.execute(
                "INSERT INTO database_backups (label, payload, created_at) VALUES (?, ?, ?)",
                (backup_name, json.dumps(snapshot, ensure_ascii=False, default=str), now_text()),
            )
        return BackupReference(backup_name)
    BACKUP_DIR.mkdir(parents=True, exist_ok=True)
    safe_label = re.sub(r"[^A-Za-z0-9_-]+", "-", label).strip("-") or "backup"
    destination = BACKUP_DIR / f"imei-pos-{datetime.now().strftime('%Y%m%d-%H%M%S-%f')}-{safe_label}.db"
    with sqlite3.connect(DB_PATH) as source, sqlite3.connect(destination) as target:
        source.backup(target)
    return destination


def add_audit(con: DatabaseConnection, aksi: str, imei: str = "", detail: Any = None) -> None:
    con.execute(
        "INSERT INTO audit_log (aksi, imei, detail, created_at) VALUES (?, ?, ?, ?)",
        (aksi, imei, json.dumps(detail, ensure_ascii=False, default=str) if detail is not None else "", now_text()),
    )


def admin_find_items(imei_input: str) -> list[dict[str, Any]]:
    imei_input = clean_imei(imei_input)
    if not imei_input:
        raise ValueError("Masukkan IMEI lengkap atau 6 angka terakhir.")
    pattern = imei_input if len(imei_input) > 6 else f"%{imei_input}"
    operator = "=" if len(imei_input) > 6 else "LIKE"
    with db() as con:
        rows = con.execute(
            f"""
            SELECT b.*, p.id AS sale_id, p.harga_jual, p.laba_rugi,
                   p.customer, p.batch_id, p.tanggal_jual
            FROM barang b
            LEFT JOIN penjualan p ON p.id = (
                SELECT p2.id FROM penjualan p2
                WHERE p2.imei = b.imei ORDER BY p2.id DESC LIMIT 1
            )
            WHERE b.imei {operator} ?
            ORDER BY b.imei
            LIMIT 20
            """,
            (pattern,),
        ).fetchall()
    return [row_to_dict(row) for row in rows if row]


def admin_update_item(payload: dict[str, Any]) -> dict[str, Any]:
    original_imei = clean_imei(payload.get("original_imei"))
    new_imei = clean_imei(payload.get("imei"))
    tipe = clean_text(payload.get("tipe"))
    hpp = to_number(payload.get("hpp"))
    if not original_imei or not new_imei or not tipe:
        raise ValueError("IMEI dan tipe wajib diisi.")
    with db() as con:
        old = con.execute("SELECT * FROM barang WHERE imei = ?", (original_imei,)).fetchone()
        if not old:
            raise ValueError("Barang tidak ditemukan.")
        if old["status"] != "BELUM_LAKU":
            raise ValueError("Barang terjual tidak dapat diedit dari data stok. Koreksi transaksi pada bagian penjualan.")
    backup = create_database_backup("sebelum-edit-barang")
    with db() as con:
        con.execute(
            """
            UPDATE barang SET imei=?, tipe=?, hpp=?, modal_awal=?, grade=?, spek=?, keterangan=?,
                              tanggal_masuk=?, supplier=?, invoice=?, updated_at=?
            WHERE imei=? AND status='BELUM_LAKU'
            """,
            (new_imei, tipe, hpp, to_number(payload.get("modal_awal")), clean_text(payload.get("grade")), clean_text(payload.get("spek")),
             clean_text(payload.get("keterangan")), clean_text(payload.get("tanggal_masuk")), clean_text(payload.get("supplier")),
             clean_text(payload.get("invoice")), now_text(), original_imei),
        )
        add_audit(con, "EDIT_BARANG", new_imei, {"sebelum": row_to_dict(old), "sesudah": payload})
    return {"message": "Data barang diperbarui.", "imei": new_imei, "backup": backup.name}


def admin_delete_item(payload: dict[str, Any]) -> dict[str, Any]:
    imei = clean_imei(payload.get("imei"))
    with db() as con:
        old = con.execute("SELECT * FROM barang WHERE imei = ?", (imei,)).fetchone()
        if not old:
            raise ValueError("Barang tidak ditemukan.")
        if old["status"] != "BELUM_LAKU":
            raise ValueError("Barang terjual harus dibatalkan transaksinya terlebih dahulu.")
    backup = create_database_backup("sebelum-hapus-barang")
    with db() as con:
        con.execute("DELETE FROM barang WHERE imei = ? AND status = 'BELUM_LAKU'", (imei,))
        add_audit(con, "HAPUS_BARANG", imei, row_to_dict(old))
    return {"message": "Barang belum terjual berhasil dihapus.", "backup": backup.name}


def admin_update_sale(payload: dict[str, Any]) -> dict[str, Any]:
    sale_id = int(payload.get("sale_id") or 0)
    harga_jual = to_number(payload.get("harga_jual"))
    alasan = clean_text(payload.get("alasan"))
    if not sale_id or harga_jual is None:
        raise ValueError("Transaksi dan harga jual wajib diisi.")
    if not alasan:
        raise ValueError("Alasan koreksi wajib diisi.")
    with db() as con:
        old = con.execute("SELECT * FROM penjualan WHERE id = ?", (sale_id,)).fetchone()
        if not old:
            raise ValueError("Transaksi tidak ditemukan.")
    backup = create_database_backup("sebelum-koreksi-penjualan")
    laba_rugi = harga_jual - old["hpp"] if old["hpp"] is not None else None
    with db() as con:
        con.execute("UPDATE penjualan SET harga_jual = ?, laba_rugi = ? WHERE id = ?", (harga_jual, laba_rugi, sale_id))
        add_audit(con, "KOREKSI_PENJUALAN", old["imei"], {"harga_lama": old["harga_jual"], "harga_baru": harga_jual, "alasan": alasan})
    return {"message": "Harga jual berhasil dikoreksi.", "laba_rugi": laba_rugi, "backup": backup.name}


def admin_void_sale(payload: dict[str, Any]) -> dict[str, Any]:
    sale_id = int(payload.get("sale_id") or 0)
    alasan = clean_text(payload.get("alasan"))
    if not sale_id or not alasan:
        raise ValueError("Transaksi dan alasan pembatalan wajib diisi.")
    with db() as con:
        sale = con.execute("SELECT * FROM penjualan WHERE id = ?", (sale_id,)).fetchone()
        if not sale:
            raise ValueError("Transaksi tidak ditemukan.")
    backup = create_database_backup("sebelum-batalkan-penjualan")
    with db() as con:
        con.execute("DELETE FROM penjualan WHERE id = ?", (sale_id,))
        remaining = con.execute("SELECT COUNT(*) c FROM penjualan WHERE imei = ?", (sale["imei"],)).fetchone()["c"]
        if not remaining:
            con.execute("UPDATE barang SET status='BELUM_LAKU', updated_at=? WHERE imei=?", (now_text(), sale["imei"]))
        add_audit(con, "BATAL_PENJUALAN", sale["imei"], {"transaksi": row_to_dict(sale), "alasan": alasan})
    return {"message": "Transaksi dibatalkan dan barang kembali menjadi stok.", "backup": backup.name}


def admin_reset_database(payload: dict[str, Any]) -> dict[str, Any]:
    if clean_text(payload.get("konfirmasi")) != "HAPUS SEMUA DATA":
        raise ValueError('Ketik tepat "HAPUS SEMUA DATA" untuk melanjutkan.')
    backup = create_database_backup("sebelum-kosongkan-database")
    with db() as con:
        counts = {
            "barang": con.execute("SELECT COUNT(*) c FROM barang").fetchone()["c"],
            "penjualan": con.execute("SELECT COUNT(*) c FROM penjualan").fetchone()["c"],
        }
        con.execute("DELETE FROM penjualan")
        con.execute("DELETE FROM barang")
        con.execute("DELETE FROM import_log")
        if con.backend == "sqlite":
            con.execute("DELETE FROM sqlite_sequence WHERE name IN ('penjualan','import_log')")
        add_audit(con, "KOSONGKAN_DATABASE", "", {"jumlah_sebelum": counts, "backup": backup.name})
    return {"message": "Database berhasil dikosongkan.", "backup": backup.name, "deleted": counts}


def sell_item(payload: dict[str, Any]) -> dict[str, Any]:
    imei_input = clean_imei(payload.get("imei"))
    harga_jual = to_number(payload.get("harga_jual"))
    customer = clean_text(payload.get("customer"))
    catatan = clean_text(payload.get("catatan"))
    if not imei_input:
        raise ValueError("IMEI wajib diisi.")
    if harga_jual is None:
        raise ValueError("Harga jual wajib diisi.")

    timestamp = now_text()
    with db() as con:
        resolved = resolve_single_item_for_sale(imei_input)
        item = con.execute("SELECT * FROM barang WHERE imei = ?", (resolved["imei"],)).fetchone()
        hpp = item["hpp"]
        laba_rugi = harga_jual - hpp if hpp is not None else None
        sale_id = con.insert_and_get_id(
            """
            INSERT INTO penjualan
            (imei, tipe, hpp, harga_jual, laba_rugi, customer, catatan, batch_id, tanggal_jual, created_at)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (item["imei"], item["tipe"], hpp, harga_jual, laba_rugi, customer, catatan, None, timestamp, timestamp),
        )
        con.execute(
            "UPDATE barang SET status = 'TERJUAL', updated_at = ? WHERE imei = ?",
            (timestamp, item["imei"]),
        )
    return {"id": sale_id, "imei": item["imei"], "harga_jual": harga_jual, "laba_rugi": laba_rugi}


def parse_bulk_sale_lines(text: str, default_price: float | None = None, require_price: bool = False) -> list[dict[str, Any]]:
    items: list[dict[str, Any]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [part.strip() for part in re.split(r"[\t,;]+", line) if part.strip()]
        if not parts:
            continue
        imei = clean_imei(parts[0])
        price = to_number(parts[1]) if len(parts) > 1 else default_price
        if not imei:
            continue
        if require_price and price is None:
            raise ValueError(f"Harga jual belum diisi untuk IMEI {imei}.")
        items.append({"imei": imei, "harga_jual": price})
    if not items:
        raise ValueError("Daftar IMEI penjualan massal masih kosong.")
    seen: set[str] = set()
    for item in items:
        if item["imei"] in seen:
            raise ValueError(f"IMEI {item['imei']} dobel di input penjualan massal.")
        seen.add(item["imei"])
    return items


def review_bulk_sale(payload: dict[str, Any]) -> dict[str, Any]:
    default_price = to_number(payload.get("harga_jual_sama"))
    items = parse_bulk_sale_lines(payload.get("items_text", ""), default_price, require_price=False)
    rows: list[dict[str, Any]] = []
    can_submit = True
    total_jual = 0.0
    total_laba = 0.0

    for entry in items:
        row = {
            "input": entry["imei"],
            "imei": "",
            "tipe": "",
            "hpp": None,
            "grade": "",
            "spek": "",
            "harga_jual": entry["harga_jual"],
            "laba_rugi": None,
            "status": "OK",
            "message": "",
            "choices": [],
        }
        matches = find_items_by_imei_input(entry["imei"], only_unsold=True, limit=10)
        if not matches:
            row["status"] = "ERROR"
            row["message"] = "Tidak ditemukan atau sudah terjual"
            can_submit = False
        elif len(matches) > 1:
            row["status"] = "PILIH_IMEI"
            row["message"] = "6 digit cocok ke lebih dari satu barang"
            row["choices"] = [{"imei": item["imei"], "tipe": item["tipe"], "hpp": item["hpp"]} for item in matches]
            can_submit = False
        else:
            item = matches[0]
            row.update({
                "imei": item["imei"],
                "tipe": item["tipe"],
                "hpp": item["hpp"],
                "grade": item["grade"],
                "spek": item["spek"],
            })
            if row["harga_jual"] is not None and row["hpp"] is not None:
                row["laba_rugi"] = row["harga_jual"] - row["hpp"]
                total_jual += row["harga_jual"]
                total_laba += row["laba_rugi"]
            else:
                can_submit = False
                row["status"] = "BUTUH_HARGA"
                row["message"] = "Isi harga jual"
        rows.append(row)

    return {"items": rows, "can_submit": can_submit, "total_jual": total_jual, "total_laba": total_laba}


def sell_bulk(payload: dict[str, Any]) -> dict[str, Any]:
    customer = clean_text(payload.get("customer"))
    catatan = clean_text(payload.get("catatan"))
    raw_items = payload.get("items")
    if isinstance(raw_items, list) and raw_items:
        items = [{"imei": clean_imei(item.get("imei")), "harga_jual": to_number(item.get("harga_jual"))} for item in raw_items]
        for item in items:
            if not item["imei"]:
                raise ValueError("Ada IMEI kosong di penjualan massal.")
            if item["harga_jual"] is None:
                raise ValueError(f"Harga jual belum diisi untuk IMEI {item['imei']}.")
    else:
        default_price = to_number(payload.get("harga_jual_sama"))
        items = parse_bulk_sale_lines(payload.get("items_text", ""), default_price, require_price=True)
    timestamp = now_text()
    batch_id = f"BT-{datetime.now().strftime('%Y%m%d%H%M%S')}"
    total_jual = 0.0
    total_laba = 0.0
    sold_rows: list[dict[str, Any]] = []

    with db() as con:
        resolved_rows: list[tuple[Any, float]] = []
        for entry in items:
            matches = find_items_by_imei_input(entry["imei"], only_unsold=True, limit=10)
            if not matches:
                raise ValueError(f"IMEI {entry['imei']} tidak ditemukan atau sudah terjual.")
            if len(matches) > 1:
                choices = ", ".join(item["imei"] for item in matches[:5])
                raise ValueError(f"Input {entry['imei']} cocok ke lebih dari 1 barang. Pakai IMEI lengkap: {choices}")
            item = con.execute("SELECT * FROM barang WHERE imei = ?", (matches[0]["imei"],)).fetchone()
            resolved_rows.append((item, entry["harga_jual"]))

        sale_ids: list[int] = []
        for item, harga_jual in resolved_rows:
            hpp = item["hpp"]
            laba_rugi = harga_jual - hpp if hpp is not None else None
            sale_id = con.insert_and_get_id(
                """
                INSERT INTO penjualan
                (imei, tipe, hpp, harga_jual, laba_rugi, customer, catatan, batch_id, tanggal_jual, created_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (item["imei"], item["tipe"], hpp, harga_jual, laba_rugi, customer, catatan, batch_id, timestamp, timestamp),
            )
            sale_ids.append(sale_id)
            con.execute("UPDATE barang SET status = 'TERJUAL', updated_at = ? WHERE imei = ?", (timestamp, item["imei"]))
            total_jual += harga_jual
            total_laba += laba_rugi or 0
            sold_rows.append({"imei": item["imei"], "tipe": item["tipe"], "harga_jual": harga_jual})

    return {
        "batch_id": batch_id,
        "ids": sale_ids,
        "jumlah": len(sold_rows),
        "total_jual": total_jual,
        "total_laba": total_laba,
        "items": sold_rows,
    }


def manual_inbound(payload: dict[str, Any]) -> dict[str, Any]:
    imei = clean_imei(payload.get("imei"))
    tipe = clean_text(payload.get("tipe"))
    hpp = to_number(payload.get("hpp"))
    if not imei or not tipe or hpp is None:
        raise ValueError("IMEI, tipe, dan HPP wajib diisi.")
    record = {
        "imei": imei,
        "tipe": tipe,
        "hpp": hpp,
        "modal_awal": to_number(payload.get("modal_awal")) or hpp,
        "grade": clean_text(payload.get("grade")),
        "spek": clean_text(payload.get("spek")),
        "keterangan": clean_text(payload.get("keterangan")),
        "sumber_file": "Input manual",
        "sumber_sheet": "Manual",
    }
    return import_records([record], "Input manual", payload)


def get_sale(sale_id: int) -> dict[str, Any] | None:
    with db() as con:
        row = con.execute("SELECT * FROM penjualan WHERE id = ?", (sale_id,)).fetchone()
    return row_to_dict(row)


def get_batch_sales(batch_id: str) -> list[dict[str, Any]]:
    with db() as con:
        rows = con.execute("SELECT * FROM penjualan WHERE batch_id = ? ORDER BY id", (batch_id,)).fetchall()
    return [row_to_dict(row) for row in rows if row]


def nota_html(sale_id: int) -> str:
    sale = get_sale(sale_id)
    if not sale:
        return "<h1>Nota tidak ditemukan</h1>"

    nomor = f"NT-{sale['id']:06d}"
    customer = sale["customer"] or "-"
    catatan = sale["catatan"] or "-"
    return f"""
<!doctype html>
<html lang="id">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Nota {nomor}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 0; background: #f5f7fb; color: #101828; }}
    .nota {{ width: 360px; margin: 24px auto; background: white; padding: 20px; border-radius: 12px; box-shadow: 0 2px 12px rgba(16,24,40,.12); }}
    h1 {{ font-size: 20px; margin: 0; text-align: center; }}
    .store-header {{ text-align: center; margin-bottom: 12px; }}
    .store-logo {{ width: 105px; height: 105px; object-fit: contain; display: block; margin: 0 auto 4px; }}
    .store-address {{ color: #475467; font-size: 11px; line-height: 1.45; }}
    .center {{ text-align: center; }}
    .muted {{ color: #667085; font-size: 12px; }}
    .line {{ border-top: 1px dashed #98a2b3; margin: 14px 0; }}
    .row {{ display: flex; justify-content: space-between; gap: 12px; margin: 7px 0; }}
    .label {{ color: #475467; }}
    .value {{ text-align: right; font-weight: 600; }}
    .total {{ font-size: 18px; font-weight: 700; }}
    button {{ width: 100%; padding: 10px; border: 0; border-radius: 10px; background: #1f4e78; color: white; font-weight: 700; cursor: pointer; margin-top: 14px; }}
    @media print {{
      body {{ background: white; }}
      .nota {{ width: auto; margin: 0; box-shadow: none; border-radius: 0; }}
      button {{ display: none; }}
    }}
  </style>
</head>
<body>
  <div class="nota">
    <div class="store-header">
      <img class="store-logo" src="/assets/find-cell-logo.png" alt="Find Cell">
      <div class="store-address">Jln Kaliurang KM 6,5 SDN Condongcatur. Yogyakarta<br>Telp. 081910055770</div>
    </div>
    <h1>NOTA PENJUALAN</h1>
    <div class="line"></div>
    <div class="row"><span class="label">No Nota</span><span class="value">{nomor}</span></div>
    <div class="row"><span class="label">Tanggal</span><span class="value">{sale['tanggal_jual']}</span></div>
    <div class="row"><span class="label">Customer</span><span class="value">{customer}</span></div>
    <div class="line"></div>
    <div class="row"><span class="label">IMEI</span><span class="value">{sale['imei']}</span></div>
    <div class="row"><span class="label">Tipe</span><span class="value">{sale['tipe']}</span></div>
    <div class="row total"><span>Harga Jual</span><span>{rupiah(sale['harga_jual'])}</span></div>
    <div class="line"></div>
    <div class="muted">Catatan: {catatan}</div>
    <p class="center muted">Terima kasih.</p>
    <button onclick="window.print()">Cetak Nota</button>
  </div>
</body>
</html>
"""


def nota_batch_html(batch_id: str) -> str:
    sales = get_batch_sales(batch_id)
    if not sales:
        return "<h1>Nota batch tidak ditemukan</h1>"
    first = sales[0]
    customer = first["customer"] or "-"
    catatan = first["catatan"] or "-"
    total = sum((sale["harga_jual"] or 0) for sale in sales)
    rows_html = "".join(
        f"""
        <tr>
          <td>{idx}</td>
          <td>{sale['imei']}</td>
          <td>{sale['tipe']}</td>
          <td class="right">{rupiah(sale['harga_jual'])}</td>
        </tr>
        """
        for idx, sale in enumerate(sales, start=1)
    )
    return f"""
<!doctype html>
<html lang="id">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Nota {batch_id}</title>
  <style>
    body {{ font-family: Arial, sans-serif; margin: 0; background: #f5f7fb; color: #101828; }}
    .nota {{ width: 720px; margin: 24px auto; background: white; padding: 22px; border-radius: 12px; box-shadow: 0 2px 12px rgba(16,24,40,.12); }}
    h1 {{ font-size: 20px; margin: 0; text-align: center; }}
    .store-header {{ text-align: center; margin-bottom: 12px; }}
    .store-logo {{ width: 105px; height: 105px; object-fit: contain; display: block; margin: 0 auto 4px; }}
    .store-address {{ color: #475467; font-size: 11px; line-height: 1.45; }}
    .center {{ text-align: center; }}
    .muted {{ color: #667085; font-size: 12px; }}
    .line {{ border-top: 1px dashed #98a2b3; margin: 14px 0; }}
    .meta {{ display: grid; grid-template-columns: 1fr 1fr; gap: 8px 24px; }}
    .row {{ display: flex; justify-content: space-between; gap: 12px; }}
    .value {{ font-weight: 600; }}
    table {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
    th, td {{ border-bottom: 1px solid #eaecf0; padding: 8px; text-align: left; }}
    th {{ background: #f2f4f7; }}
    .right {{ text-align: right; }}
    .total {{ font-size: 18px; font-weight: 700; }}
    button {{ width: 100%; padding: 10px; border: 0; border-radius: 10px; background: #1f4e78; color: white; font-weight: 700; cursor: pointer; margin-top: 14px; }}
    @media print {{
      body {{ background: white; }}
      .nota {{ width: auto; margin: 0; box-shadow: none; border-radius: 0; }}
      button {{ display: none; }}
    }}
  </style>
</head>
<body>
  <div class="nota">
    <div class="store-header">
      <img class="store-logo" src="/assets/find-cell-logo.png" alt="Find Cell">
      <div class="store-address">Jln Kaliurang KM 6,5 SDN Condongcatur. Yogyakarta<br>Telp. 081910055770</div>
    </div>
    <h1>NOTA PENJUALAN MASSAL</h1>
    <div class="line"></div>
    <div class="meta">
      <div class="row"><span>No Batch</span><span class="value">{batch_id}</span></div>
      <div class="row"><span>Tanggal</span><span class="value">{first['tanggal_jual']}</span></div>
      <div class="row"><span>Customer</span><span class="value">{customer}</span></div>
      <div class="row"><span>Jumlah</span><span class="value">{len(sales)} pcs</span></div>
    </div>
    <div class="line"></div>
    <table>
      <tr><th>No</th><th>IMEI</th><th>Tipe</th><th class="right">Harga jual</th></tr>
      {rows_html}
    </table>
    <div class="line"></div>
    <div class="row total"><span>Total</span><span>{rupiah(total)}</span></div>
    <div class="line"></div>
    <div class="muted">Catatan: {catatan}</div>
    <p class="center muted">Terima kasih.</p>
    <button onclick="window.print()">Cetak Nota</button>
  </div>
</body>
</html>
"""


def parse_text_rows(text: str, source_name: str = "Input manual") -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if "IMEI" in line.upper() and ("TIPE" in line.upper() or "TYPE" in line.upper()):
            continue
        if "\t" in line:
            parts = [part.strip() for part in line.split("\t")]
        else:
            parts = [part.strip() for part in line.split(",")]
        if len(parts) < 3:
            continue
        imei, tipe, hpp = parts[0], parts[1], parts[2]
        record = {
            "imei": clean_imei(imei),
            "tipe": clean_text(tipe),
            "hpp": to_number(hpp),
            "modal_awal": to_number(hpp),
            "grade": clean_text(parts[3]) if len(parts) > 3 else "",
            "spek": clean_text(parts[4]) if len(parts) > 4 else "",
            "keterangan": clean_text(parts[5]) if len(parts) > 5 else "",
            "sumber_file": source_name,
            "sumber_sheet": "Paste manual",
        }
        if record["imei"] and record["tipe"]:
            records.append(record)
    deduped: dict[str, dict[str, Any]] = {}
    for record in records:
        deduped[record["imei"]] = record
    return list(deduped.values())


def export_excel(start_date: str = "", end_date: str = "") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ringkasan"
    s = summary(start_date, end_date)
    ws.append(["Ringkasan Sistem Stok IMEI"])
    ws.append(["Dibuat", now_text()])
    ws.append([])
    ws.append(["Stok belum laku", s["stok"]])
    ws.append(["Barang terjual", s["terjual"]])
    ws.append(["Total penjualan", s["total_jual"]])
    ws.append(["Total laba/rugi", s["total_laba"]])
    ws.append(["Nilai modal stok", s["nilai_stok"]])
    ws.append(["Periode", f"{start_date or 'awal'} s.d. {end_date or 'sekarang'}"])

    stok_ws = wb.create_sheet("Stok")
    stok_headers = ["IMEI", "Tipe", "HPP", "Modal awal", "Grade", "Spek", "Keterangan", "Status", "Tanggal masuk", "Sumber file", "Sheet"]
    stok_ws.append(stok_headers)
    with db() as con:
        for row in con.execute("SELECT * FROM barang ORDER BY status, tipe, imei"):
            stok_ws.append([
                row["imei"], row["tipe"], row["hpp"], row["modal_awal"], row["grade"], row["spek"],
                row["keterangan"], row["status"], row["tanggal_masuk"], row["sumber_file"], row["sumber_sheet"],
            ])

    jual_ws = wb.create_sheet("Penjualan")
    jual_headers = ["Tanggal jual", "IMEI", "Tipe", "HPP", "Harga jual", "Laba/rugi", "Customer", "Catatan"]
    jual_ws.append(jual_headers)
    for row in list_penjualan(100000, start_date, end_date):
        jual_ws.append([
            row["tanggal_jual"], row["imei"], row["tipe"], row["hpp"], row["harga_jual"],
            row["laba_rugi"], row["customer"], row["catatan"],
        ])

    masuk_ws = wb.create_sheet("Barang Masuk")
    masuk_ws.append(["Batch", "Tanggal pembelian", "Supplier", "Invoice", "Jumlah unit", "Total modal", "Sumber"])
    for row in list_purchase_batches(100000):
        masuk_ws.append([row["batch_id"], row["tanggal_pembelian"], row["supplier"], row["invoice"], row["jumlah_baris"], row["total_modal"], row["source_name"]])

    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center")
        for col in range(1, sheet.max_column + 1):
            letter = get_column_letter(col)
            sheet.column_dimensions[letter].width = min(35, max(12, max(len(clean_text(sheet.cell(row=r, column=col).value)) for r in range(1, min(sheet.max_row, 200) + 1)) + 2))
        for row in sheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0'
        sheet.freeze_panes = "A2"

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


HTML = r"""
<!doctype html>
<html lang="id">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Find Cell - Sistem Penjualan</title>
  <style>
    :root { --blue:#1f5a8a; --bg:#eef3f8; --card:#fff; --soft:#f7fafc; --text:#172033; --muted:#667085; --line:#d9e2ec; --green:#15803d; --red:#b42318; }
    * { box-sizing:border-box; }
    body { margin:0; font-family:Arial,sans-serif; background:var(--bg); color:var(--text); }
    .topbar { height:66px; display:flex; align-items:center; gap:16px; padding:12px 18px; background:var(--card); border-bottom:1px solid var(--line); position:sticky; top:0; z-index:10; }
    .brand { display:flex; align-items:center; gap:10px; font-weight:800; white-space:nowrap; }
    .app-logo { width:46px; height:46px; border-radius:10px; object-fit:contain; background:white; }
    .brand-subtitle { color:var(--muted); font-size:11px; font-weight:500; margin-top:2px; }
    .scan-top { flex:1; max-width:520px; display:flex; gap:8px; }
    input, textarea, button, select { font:inherit; }
    input, textarea, select { width:100%; border:1px solid var(--line); background:var(--soft); padding:11px 12px; border-radius:12px; outline:none; }
    textarea { resize:vertical; }
    button { border:0; border-radius:12px; padding:11px 14px; background:var(--blue); color:white; font-weight:700; cursor:pointer; }
    button.secondary { background:#475467; }
    button.ghost { background:var(--soft); color:var(--text); border:1px solid var(--line); }
    button.green { background:var(--green); }
    button.danger { background:var(--red); }
    button:disabled { background:#d0d5dd !important; color:#667085 !important; cursor:not-allowed; box-shadow:none; opacity:1; }
    .layout { display:grid; grid-template-columns:220px minmax(0,1fr); min-height:calc(100vh - 66px); }
    .side { background:var(--card); border-right:1px solid var(--line); padding:16px; }
    .nav-title { color:var(--muted); font-size:12px; text-transform:uppercase; letter-spacing:.04em; margin:18px 0 8px; }
    .nav-btn { width:100%; background:transparent; color:var(--text); text-align:left; margin:2px 0; padding:10px 12px; }
    .nav-btn.active { background:var(--blue); color:white; }
    main { min-width:0; padding:18px; }
    .page { display:none; }
    .page.active { display:block; }
    .grid { display:grid; grid-template-columns:repeat(auto-fit,minmax(230px,1fr)); gap:14px; }
    .hero { display:grid; grid-template-columns:1.2fr .8fr; gap:14px; margin-bottom:14px; }
    .card { background:var(--card); border:1px solid var(--line); border-radius:16px; padding:16px; }
    h1, h2, h3 { margin:0 0 10px; }
    h1 { font-size:22px; }
    h2 { font-size:18px; }
    h3 { font-size:15px; }
    .muted { color:var(--muted); font-size:13px; }
    .ok { color:var(--green); font-weight:700; }
    .bad { color:var(--red); font-weight:700; }
    .stat strong { display:block; font-size:26px; margin-top:4px; }
    .row { display:grid; grid-template-columns:1fr 1fr; gap:10px; }
    .report-card { padding:18px; }
    .report-header { align-items:flex-start; margin-bottom:16px; }
    .report-header h1 { margin-bottom:4px; }
    .report-download { white-space:nowrap; }
    .report-controls { padding-top:14px; border-top:1px solid #eaecf0; }
    .report-control-main { min-width:0; }
    .report-mode { display:inline-flex; padding:4px; border:1px solid var(--line); border-radius:12px; background:var(--soft); margin-bottom:12px; }
    .report-mode button { min-width:132px; padding:8px 12px; background:transparent; color:var(--muted); }
    .report-mode button.active { background:var(--card); color:var(--blue); border:1px solid var(--line); box-shadow:0 1px 2px rgba(16,24,40,.06); }
    .report-filters { display:grid; grid-template-columns:1fr 1fr; gap:10px; align-items:end; }
    .report-filters.month-mode { grid-template-columns:minmax(220px,420px); }
    .field-label { display:grid; gap:6px; color:var(--muted); font-size:13px; }
    .report-period { display:inline-flex; align-items:center; gap:7px; margin-top:12px; padding:7px 10px; border-radius:999px; background:#edf6ff; color:var(--blue); font-size:12px; font-weight:700; }
    .report-period::before { content:''; width:7px; height:7px; border-radius:50%; background:var(--blue); }
    .report-stat { position:relative; overflow:hidden; min-height:108px; display:flex; flex-direction:column; justify-content:center; padding:18px 20px; border-left:4px solid var(--blue); }
    .report-stat.units { border-left-color:#0f8a5f; }
    .report-stat strong { font-size:30px; line-height:1.15; }
    .report-section-grid { margin-top:14px; }
    .ranking-card h2 { font-size:16px; margin-bottom:4px; }
    .ranking-list { display:grid; gap:0; }
    .ranking-row { width:100%; display:grid; grid-template-columns:32px minmax(0,1fr) auto; gap:10px; align-items:center; padding:10px 6px; border:0; border-bottom:1px solid #eaecf0; border-radius:8px; background:transparent; color:var(--text); text-align:left; }
    .ranking-row:last-child { border-bottom:0; }
    .ranking-row:hover { background:var(--soft); }
    .ranking-row.active { background:#edf6ff; box-shadow:inset 0 0 0 1px #c9def1; }
    .ranking-number { width:27px; height:27px; display:grid; place-items:center; border-radius:999px; background:#e9f2fb; color:var(--blue); font-weight:800; font-size:12px; }
    .ranking-name { font-weight:700; }
    .ranking-meta { display:block; color:var(--muted); font-size:12px; font-weight:400; margin-top:2px; }
    .ranking-count { font-weight:800; white-space:nowrap; }
    .empty-state { color:var(--muted); padding:14px 0; }
    .report-table-card { margin-top:14px; }
    .report-table-tools { display:grid; grid-template-columns:minmax(260px,1fr) auto; gap:12px; align-items:center; margin:12px 0; }
    .report-search-area { display:flex; gap:8px; align-items:center; min-width:0; }
    .report-search-area input { max-width:560px; }
    .report-filter-chip { display:inline-flex; align-items:center; gap:7px; padding:7px 9px; border-radius:999px; background:#edf6ff; color:var(--blue); font-size:12px; font-weight:700; white-space:nowrap; }
    .report-filter-chip button { width:22px; height:22px; padding:0; border-radius:50%; background:var(--blue); display:grid; place-items:center; }
    .report-table-meta { display:grid; justify-items:end; gap:3px; color:var(--muted); font-size:12px; white-space:nowrap; }
    .report-result-count { color:var(--text); font-weight:700; }
    .report-status.bad { color:var(--red); }
    .report-table-wrap { border:1px solid #eaecf0; border-radius:12px; max-height:460px; }
    .report-table { min-width:620px; }
    .report-table th { position:sticky; top:0; z-index:1; padding:11px 12px; color:#475467; font-size:12px; text-transform:uppercase; letter-spacing:.02em; }
    .report-table td { padding:11px 12px; }
    .report-table tr:nth-child(odd):not(:first-child) td { background:#fbfcfd; }
    .report-table tr:hover td { background:#f2f7fb; }
    .report-detail-button { width:auto; padding:7px 10px; background:var(--soft); color:var(--blue); border:1px solid var(--line); }
    .report-detail-row td { background:#f8fbfe !important; padding:0 12px 12px; }
    .report-detail-grid { display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:10px; padding:12px; border:1px solid #dce8f2; border-radius:10px; background:white; }
    .report-detail-item { min-width:0; }
    .report-detail-label { color:var(--muted); font-size:11px; text-transform:uppercase; letter-spacing:.03em; margin-bottom:3px; }
    .report-detail-value { font-weight:700; overflow-wrap:anywhere; }
    .cell-right { text-align:right; white-space:nowrap; }
    .cell-nowrap { white-space:nowrap; }
    .hidden { display:none !important; }
    .entry-tabs { display:flex; gap:8px; flex-wrap:wrap; margin:12px 0 14px; }
    .entry-tabs button { background:var(--soft); color:var(--text); border:1px solid var(--line); }
    .entry-tabs button.active { background:var(--blue); color:white; border-color:var(--blue); }
    .entry-panel { margin-top:0; }
    .inventory-filters, .sales-history-filters { display:grid; grid-template-columns:minmax(220px,1.4fr) minmax(160px,.8fr) minmax(160px,.8fr); gap:10px; align-items:end; }
    .compact-toolbar { display:flex; justify-content:space-between; align-items:center; gap:10px; margin:12px 0 8px; }
    .compact-count { color:var(--muted); font-size:13px; font-weight:700; }
    .compact-table { min-width:620px; }
    .status-badge { display:inline-flex; align-items:center; padding:5px 9px; border-radius:999px; font-size:12px; font-weight:800; white-space:nowrap; }
    .status-available { background:#eaf7ef; color:#067647; }
    .status-sold { background:#f2f4f7; color:#475467; }
    .advanced-panel { margin-top:14px; background:var(--card); border:1px solid var(--line); border-radius:16px; overflow:hidden; }
    .advanced-panel summary { cursor:pointer; padding:15px 16px; font-weight:800; list-style:none; }
    .advanced-panel summary::-webkit-details-marker { display:none; }
    .advanced-panel summary::after { content:'+'; float:right; color:var(--blue); font-size:20px; line-height:16px; }
    .advanced-panel[open] summary::after { content:'−'; }
    .advanced-content { border-top:1px solid var(--line); padding:16px; }
    .scan-box { display:flex; gap:10px; margin-top:12px; }
    .split { display:grid; grid-template-columns:.9fr 1.1fr; gap:14px; margin-top:14px; }
    .section-head { display:flex; align-items:center; justify-content:space-between; gap:10px; margin-bottom:10px; }
    .switch { display:flex; background:var(--soft); border:1px solid var(--line); border-radius:12px; padding:4px; }
    .switch button { background:transparent; color:var(--muted); padding:7px 10px; }
    .switch button.active { background:var(--card); color:var(--text); border:1px solid var(--line); }
    .group-list, .detail-list { display:grid; gap:9px; }
    .group-btn { background:var(--soft); color:var(--text); border:1px solid var(--line); text-align:left; display:flex; align-items:center; justify-content:space-between; gap:10px; }
    .group-btn.active { border-color:var(--blue); background:#edf6ff; }
    .pill { display:inline-block; border-radius:999px; background:#e9f2fb; color:var(--blue); padding:5px 9px; font-size:12px; font-weight:700; }
    .detail-row { display:grid; grid-template-columns:1fr auto; gap:10px; padding:11px 0; border-bottom:1px solid var(--line); }
    .detail-row:last-child { border-bottom:0; }
    .amount { font-weight:800; }
    .upload-grid { display:grid; grid-template-columns:1fr 1fr; gap:14px; }
    .upload-card { background:var(--soft); border:1px solid var(--line); border-radius:14px; padding:14px; display:grid; gap:10px; }
    .table-wrap { overflow:auto; max-height:420px; }
    table { width:100%; border-collapse:collapse; font-size:13px; }
    th, td { border-bottom:1px solid #eaecf0; padding:9px; text-align:left; vertical-align:top; }
    th { background:var(--soft); }
    .item-choice { border:1px solid var(--line); background:var(--soft); border-radius:12px; padding:10px; margin:8px 0; }
    .modal-backdrop { display:none; position:fixed; inset:0; background:rgba(15,23,32,.48); z-index:50; padding:24px; }
    .modal-backdrop.active { display:grid; place-items:center; }
    .modal { background:var(--card); border-radius:16px; border:1px solid var(--line); width:min(860px,100%); max-height:82vh; overflow:auto; padding:16px; box-shadow:0 20px 50px rgba(16,24,40,.25); }
    .modal-head { display:flex; justify-content:space-between; gap:12px; align-items:center; margin-bottom:10px; }
    .modal-close { width:auto; background:var(--soft); color:var(--text); border:1px solid var(--line); }
    @media (max-width:820px) {
      body { overflow-x:hidden; }
      .layout { display:block; min-height:0; }
      .hero,.split,.upload-grid,.report-filters,.report-table-tools,.report-detail-grid,.inventory-filters,.sales-history-filters { grid-template-columns:1fr; }
      .topbar { height:auto; flex-direction:column; align-items:stretch; gap:10px; padding:10px 12px; }
      .scan-top { max-width:none; min-width:0; }
      .scan-top input { min-width:0; }
      .scan-top button { flex:0 0 auto; }
      .side { display:grid; grid-template-columns:repeat(3,minmax(0,1fr)); gap:6px; overflow:visible; padding:8px 12px; border-right:0; border-bottom:1px solid var(--line); }
      .nav-btn { width:100%; min-height:44px; margin:0; padding:8px 6px; text-align:center; white-space:normal; line-height:1.2; font-size:13px; }
      main { padding:10px; }
      .row { grid-template-columns:1fr; }
      .section-head { flex-wrap:wrap; align-items:flex-start; }
      .report-header { flex-direction:column; align-items:stretch; gap:12px; }
      .report-header > div { width:100%; }
      .report-download { width:100%; }
      .report-mode { display:flex; width:100%; }
      .report-mode button { min-width:0; flex:1; }
      .report-search-area { align-items:stretch; flex-direction:column; }
      .report-table-meta { justify-items:start; }
      .report-table-wrap,.table-wrap { max-width:100%; overscroll-behavior-inline:contain; -webkit-overflow-scrolling:touch; }
      .compact-toolbar { align-items:stretch; flex-direction:column; }
      .modal-backdrop { padding:10px; }
      .modal { max-height:90vh; padding:14px; }
    }
    @media (max-width:480px) {
      .app-logo { width:40px; height:40px; }
      .brand { gap:8px; }
      .scan-top { gap:6px; }
      .scan-top input { padding:10px; }
      .scan-top button { padding:10px 12px; }
      .side { grid-template-columns:repeat(2,minmax(0,1fr)); }
      .nav-btn:last-child { grid-column:1 / -1; }
      .card,.report-card { padding:14px; border-radius:14px; }
      h1 { font-size:21px; }
      h2 { font-size:17px; }
      .report-stat { min-height:92px; padding:15px 16px; }
      .report-stat strong { font-size:27px; }
      .ranking-row { grid-template-columns:28px minmax(0,1fr) auto; gap:8px; padding:10px 2px; }
      .ranking-number { width:25px; height:25px; }
      .report-table-card { padding:12px; }
      .section-head:not(.report-header) { flex-direction:column; align-items:stretch; }
      .section-head > .row { width:100%; max-width:none !important; }
      .entry-tabs { display:grid; grid-template-columns:1fr; }
      .entry-tabs button { width:100%; }
      .scan-box { flex-direction:column; }
      .scan-box button { width:100%; }
    }
  </style>
</head>
<body>
  <div class="topbar">
    <div class="brand"><img class="app-logo" src="/assets/find-cell-logo.png" alt="Find Cell"><div>Find Cell<div class="brand-subtitle">Sistem Penjualan IMEI</div></div></div>
    <div class="scan-top">
      <input id="scanImei" placeholder="Cek IMEI di seluruh data" onkeydown="if(event.key==='Enter') cekImei()">
      <button onclick="cekImei()">Cek IMEI</button>
    </div>
  </div>

  <div class="layout">
    <aside class="side">
      <button class="nav-btn active" data-page="upload" onclick="showPage('upload'); loadInboundHistory()">Barang Masuk</button>
      <button class="nav-btn" data-page="jual" onclick="showPage('jual')">Penjualan</button>
      <button class="nav-btn" data-page="persediaan" onclick="showPage('persediaan'); loadInventory()">Persediaan</button>
      <button class="nav-btn" data-page="laporan" onclick="openReport()">Laporan</button>
      <button class="nav-btn" data-page="kelola" onclick="showPage('kelola')">Master Data</button>
    </aside>

    <main>
      <section id="page-upload" class="page active">
        <div class="section-head">
          <h1>Barang masuk</h1>
          <div class="row" style="max-width:420px">
            <button class="secondary" onclick="location.href='/template/xlsx'">Template Excel</button>
            <button class="secondary" onclick="location.href='/template/csv'">Template CSV</button>
          </div>
        </div>
        <div class="entry-tabs" aria-label="Cara memasukkan barang">
          <button id="inboundModeFile" class="active" onclick="setInboundMode('file')">Upload file</button>
          <button id="inboundModePaste" onclick="setInboundMode('paste')">Paste data</button>
          <button id="inboundModeManual" onclick="setInboundMode('manual')">Input satu barang</button>
        </div>
        <div id="inboundFilePanel" class="entry-panel">
          <div class="card">
            <h2>Upload data masuk</h2>
            <div class="muted">Untuk stok awal atau impor biasa. Semua IMEI akan tercatat sebagai stok tersedia. HPP diambil dari kolom <b>Modal Pot</b>.</div>
            <div class="row" style="margin-top:10px"><input id="inSupplier" placeholder="Supplier"><input id="inInvoice" placeholder="Nomor invoice"></div>
            <input id="inDate" type="date" style="margin-top:10px">
            <div class="upload-card">
              <input id="fileDataMasuk" type="file" accept=".xlsx,.xlsm,.csv">
              <button onclick="uploadFile('DATA_MASUK')">Submit data masuk</button>
              <p id="resultDataMasuk" class="muted"></p>
            </div>
          </div>
        </div>
        <div id="inboundPastePanel" class="card entry-panel hidden">
          <h2>Upload cepat tanpa file</h2>
          <div class="muted">Paste dari Excel/CSV. Urutan kolom: IMEI, Tipe, HPP, Grade, Spek, Keterangan.</div>
          <textarea id="pasteData" rows="5" placeholder="356789123456789	iPhone 11 64GB	3500000	A	RESMI	Fullset"></textarea>
          <button onclick="uploadPaste()">Upload data paste</button>
          <p id="pasteResult" class="muted"></p>
        </div>
        <div id="inboundManualPanel" class="card entry-panel hidden">
          <h2>Input satu barang</h2>
          <div class="row"><input id="manualImei" placeholder="IMEI"><input id="manualTipe" placeholder="Tipe"></div>
          <div class="row" style="margin-top:10px"><input id="manualHpp" placeholder="HPP / Modal Pot"><input id="manualGrade" placeholder="Grade"></div>
          <div class="row" style="margin-top:10px"><input id="manualSpek" placeholder="Spek"><input id="manualKet" placeholder="Keterangan"></div>
          <button style="margin-top:10px" onclick="saveManualInbound()">Simpan barang masuk</button>
          <p id="manualInboundResult" class="muted"></p>
        </div>
        <div class="card" style="margin-top:14px">
          <div class="section-head"><h2>Riwayat barang masuk</h2><button class="ghost" onclick="loadInboundHistory()">Muat ulang</button></div>
          <div class="table-wrap"><table id="inboundTable"></table></div>
        </div>
      </section>

      <section id="page-jual" class="page">
        <div class="card">
          <h1>Jual satuan</h1>
          <div class="row" style="margin-bottom:12px"><button class="ghost" onclick="showPage('penjualan-massal')">Penjualan massal</button><button class="ghost" onclick="showPage('riwayat-penjualan'); loadSalesHistory()">Riwayat penjualan</button></div>
          <div class="row" style="margin-top:10px">
            <input id="jualImei" placeholder="IMEI lengkap atau 6 angka terakhir" onkeydown="if(event.key==='Enter') cekJualSatuan()">
            <button onclick="cekJualSatuan()">Cek detail</button>
          </div>
          <div id="jualDetail" class="upload-card" style="display:none;margin-top:10px"></div>
          <div class="row" style="margin-top:10px">
            <input id="hargaJual" placeholder="Harga jual">
            <input id="customer" placeholder="Customer / pembeli">
          </div>
          <textarea id="catatan" placeholder="Catatan transaksi" rows="2" style="margin-top:10px"></textarea>
          <button class="green" onclick="jualBarang()">Simpan jual</button>
          <p id="jualResult" class="muted"></p>
        </div>
      </section>

      <section id="page-penjualan-massal" class="page">
        <div class="card">
          <h1>Penjualan massal</h1>
          <div class="row" style="margin-bottom:12px"><button class="ghost" onclick="showPage('jual')">Jual satuan</button><button class="ghost" onclick="showPage('riwayat-penjualan'); loadSalesHistory()">Riwayat penjualan</button></div>
          <div class="muted">Sama seperti jual satuan, hanya IMEI-nya banyak. Input 6 angka terakhir atau IMEI lengkap, klik Review, lalu isi harga jual per baris sebelum submit final.</div>
          <div class="row" style="margin-top:10px">
            <input id="bulkCustomer" placeholder="Customer / pembeli">
            <input id="bulkHargaSama" placeholder="Harga jual sama, opsional">
          </div>
          <textarea id="bulkImeis" rows="7" placeholder="Input satu IMEI per baris:
261396
456789
123456" style="margin-top:10px"></textarea>
          <textarea id="bulkCatatan" rows="2" placeholder="Catatan nota massal" style="margin-top:10px"></textarea>
          <div class="row" style="margin-top:10px">
            <button onclick="reviewMassal()">Review barang</button>
            <button id="bulkSubmitBtn" class="green" onclick="jualMassal()" disabled>Submit final</button>
          </div>
          <p id="bulkResult" class="muted"></p>
          <div id="bulkReview" class="table-wrap" style="margin-top:12px"></div>
        </div>
      </section>

      <section id="page-riwayat-penjualan" class="page">
        <div class="card">
          <div class="section-head"><h1>Riwayat penjualan</h1><button class="ghost" onclick="loadSalesHistory()">Muat ulang</button></div>
          <div class="row" style="margin-bottom:12px"><button class="ghost" onclick="showPage('jual')">Jual satuan</button><button class="ghost" onclick="showPage('penjualan-massal')">Penjualan massal</button></div>
          <div class="sales-history-filters">
            <label class="field-label">Pencarian<input id="salesQuery" placeholder="IMEI, tipe, customer, atau nota" oninput="scheduleSalesHistory()" onkeydown="if(event.key==='Enter') loadSalesHistory()"></label>
            <label class="field-label">Tanggal awal<input id="salesFrom" type="date" onchange="loadSalesHistory()"></label>
            <label class="field-label">Tanggal akhir<input id="salesTo" type="date" onchange="loadSalesHistory()"></label>
          </div>
          <div class="compact-toolbar"><span id="salesHistoryCount" class="compact-count">0 transaksi</span><button onclick="loadSalesHistory()">Cari transaksi</button></div>
          <div class="table-wrap report-table-wrap"><table id="salesHistoryTable" class="report-table compact-table"></table></div>
        </div>
      </section>

      <section id="page-persediaan" class="page">
        <div class="card">
          <div class="section-head"><h1>Persediaan</h1><button class="ghost" onclick="loadInventory()">Muat ulang</button></div>
          <div class="inventory-filters">
            <label class="field-label">Pencarian<input id="inventoryQuery" placeholder="IMEI atau tipe" oninput="loadInventory()" onkeydown="if(event.key==='Enter') loadInventory()"></label>
            <label class="field-label">Status<select id="inventoryStatus" onchange="loadInventory()"><option value="">Semua status</option><option value="BELUM_LAKU">Tersedia</option><option value="TERJUAL">Terjual</option></select></label>
            <label class="field-label">Merek<select id="inventoryBrand" onchange="loadInventory()"><option value="">Semua merek</option></select></label>
          </div>
          <div class="compact-toolbar"><span id="inventoryCount" class="compact-count">0 barang</span></div>
          <div class="table-wrap report-table-wrap"><table id="inventoryTable" class="report-table compact-table"></table></div>
        </div>
        <div class="card" style="margin-top:14px">
          <h2>Riwayat IMEI</h2>
          <div class="row"><input id="historyImei" placeholder="IMEI lengkap atau 6 angka terakhir" onkeydown="if(event.key==='Enter') loadImeiHistory()"><button onclick="loadImeiHistory()">Lihat riwayat</button></div>
          <div id="imeiHistoryResult" style="margin-top:12px"></div>
        </div>
      </section>

      <section id="page-laporan" class="page">
        <div class="card report-card" style="margin-bottom:14px">
          <div class="section-head report-header">
            <div><h1>Laporan</h1><div class="muted">Ringkasan penjualan langsung berdasarkan periode yang dipilih.</div></div>
            <button class="report-download" onclick="downloadReport()">Download laporan Excel</button>
          </div>
          <div class="report-controls">
            <div class="report-control-main">
              <div class="report-mode" aria-label="Cara memilih periode">
                <button id="reportModeRange" class="active" onclick="setReportMode('range')">Rentang tanggal</button>
                <button id="reportModeMonth" onclick="setReportMode('month')">Bulanan</button>
              </div>
              <div id="reportRangeFields" class="report-filters">
                <label class="field-label">Tanggal awal<input id="reportFrom" type="date" onchange="useCustomReportDates()"></label>
                <label class="field-label">Tanggal akhir<input id="reportTo" type="date" onchange="useCustomReportDates()"></label>
              </div>
              <div id="reportMonthField" class="report-filters month-mode hidden">
                <label class="field-label">Pilih bulan<input id="reportMonth" type="month" onchange="setReportMonth()"></label>
              </div>
              <div id="reportPeriodLabel" class="report-period">Periode aktif</div>
            </div>
          </div>
        </div>
        <div class="grid">
          <div class="card stat report-stat"><span class="muted">Total pendapatan</span><strong id="reportSales">Rp 0</strong></div>
          <div class="card stat report-stat units"><span class="muted">Total barang keluar</span><strong id="reportUnits">0 unit</strong></div>
        </div>
        <div class="grid report-section-grid">
          <div class="card ranking-card"><h2>Merek paling laku</h2><div id="reportTopBrands" class="ranking-list"></div></div>
          <div class="card ranking-card"><h2>Tipe paling laku</h2><div id="reportTopTypes" class="ranking-list"></div></div>
        </div>
        <div class="card report-table-card">
          <div class="section-head"><h2>Transaksi pada periode</h2><button class="ghost" onclick="refreshReportDashboard()">Muat ulang</button></div>
          <div class="report-table-tools">
            <div class="report-search-area">
              <input id="reportSearch" placeholder="Cari nota, IMEI, tipe, merek, atau customer" oninput="renderReportTransactions()">
              <div id="reportActiveFilter" class="report-filter-chip hidden"></div>
            </div>
            <div class="report-table-meta">
              <span id="reportResultCount" class="report-result-count">0 transaksi</span>
              <span id="reportStatus" class="report-status" aria-live="polite"></span>
            </div>
          </div>
          <div class="table-wrap report-table-wrap"><table id="reportSalesTable" class="report-table"></table></div>
        </div>
      </section>

      <section id="page-kelola" class="page">
        <div class="card">
          <h1>Master Data</h1>
          <div class="muted">Cari menggunakan IMEI lengkap atau 6 angka terakhir. Perubahan penting selalu dibuatkan backup otomatis.</div>
          <div class="row" style="margin-top:12px">
            <input id="adminImei" placeholder="IMEI lengkap atau 6 angka terakhir" onkeydown="if(event.key==='Enter') adminSearch()">
            <button onclick="adminSearch()">Cari data</button>
          </div>
          <div id="adminResult" style="margin-top:12px"></div>
        </div>

        <details class="advanced-panel">
          <summary>Pengaturan lanjutan</summary>
          <div class="advanced-content">
            <h2 class="bad">Kosongkan database</h2>
            <div class="muted">Menghapus seluruh stok, penjualan, dan riwayat upload. Sistem membuat backup database terlebih dahulu.</div>
            <div class="row" style="margin-top:12px">
              <input id="resetConfirmation" placeholder='Ketik: HAPUS SEMUA DATA'>
              <button class="danger" onclick="resetDatabase()">Kosongkan database</button>
            </div>
            <p id="resetResult" class="muted"></p>
          </div>
        </details>
      </section>
    </main>
  </div>

  <div id="imeiModal" class="modal-backdrop" onclick="if(event.target.id==='imeiModal') closeModal()">
    <div class="modal">
      <div class="modal-head">
        <div>
          <h2 id="modalTitle">Daftar IMEI</h2>
          <div id="modalSubtitle" class="muted"></div>
        </div>
        <button class="modal-close" onclick="closeModal()">Tutup</button>
      </div>
      <div id="modalBody" class="table-wrap"></div>
    </div>
  </div>

  <script>
    const rupiah = (n) => n == null ? '-' : new Intl.NumberFormat('id-ID', { style:'currency', currency:'IDR', maximumFractionDigits:0 }).format(n);
    let stockStatus = 'BELUM_LAKU';
    let selectedBrand = '';
    let bulkReviewItems = [];
    let jualSatuanItem = null;
    let reportSalesRows = [];
    let reportTopBrandsRows = [];
    let reportTopTypesRows = [];
    let reportRankingFilter = {kind:'', value:''};
    let reportStatusTimer = null;
    let salesHistoryTimer = null;

    async function api(path, options={}) {
      const res = await fetch(path, options);
      const data = await res.json();
      if (!res.ok || data.error) throw new Error(data.error || 'Request gagal');
      return data;
    }
    function showPage(name) {
      document.querySelectorAll('.page').forEach(p => p.classList.remove('active'));
      document.getElementById('page-' + name).classList.add('active');
      const navName = ['jual','penjualan-massal','riwayat-penjualan'].includes(name) ? 'jual' : name;
      document.querySelectorAll('.nav-btn').forEach(b => b.classList.toggle('active', b.dataset.page === navName));
    }
    function setInboundMode(mode) {
      const modes = ['file', 'paste', 'manual'];
      modes.forEach(item => {
        const title = item.charAt(0).toUpperCase() + item.slice(1);
        document.getElementById('inboundMode' + title).classList.toggle('active', item === mode);
        document.getElementById('inbound' + title + 'Panel').classList.toggle('hidden', item !== mode);
      });
    }
    const localDateText = (date) => {
      const year = date.getFullYear();
      const month = String(date.getMonth() + 1).padStart(2, '0');
      const day = String(date.getDate()).padStart(2, '0');
      return `${year}-${month}-${day}`;
    };
    const formatReportDate = (value) => {
      if (!value) return '';
      const [year, month, day] = value.split('-').map(Number);
      return new Intl.DateTimeFormat('id-ID', {day:'numeric', month:'short', year:'numeric'}).format(new Date(year, month - 1, day));
    };
    const formatReportDateTime = (value) => {
      if (!value) return '-';
      const parsed = new Date(value.replace(' ', 'T'));
      if (Number.isNaN(parsed.getTime())) return value;
      return new Intl.DateTimeFormat('id-ID', {day:'2-digit', month:'short', year:'numeric', hour:'2-digit', minute:'2-digit'}).format(parsed);
    };
    function updateReportPeriodLabel() {
      const from = reportFrom.value;
      const to = reportTo.value;
      if (!from && !to) reportPeriodLabel.textContent = 'Semua data';
      else if (from && to && from === to) reportPeriodLabel.textContent = formatReportDate(from);
      else if (from && to) reportPeriodLabel.textContent = `${formatReportDate(from)} – ${formatReportDate(to)}`;
      else if (from) reportPeriodLabel.textContent = `Sejak ${formatReportDate(from)}`;
      else reportPeriodLabel.textContent = `Sampai ${formatReportDate(to)}`;
    }
    function setReportMode(mode) {
      const isMonth = mode === 'month';
      reportModeRange.classList.toggle('active', !isMonth);
      reportModeMonth.classList.toggle('active', isMonth);
      reportRangeFields.classList.toggle('hidden', isMonth);
      reportMonthField.classList.toggle('hidden', !isMonth);
      if (isMonth && !reportMonth.value) {
        const today = new Date();
        reportMonth.value = `${today.getFullYear()}-${String(today.getMonth() + 1).padStart(2, '0')}`;
        setReportMonth();
      }
    }
    function openReport() {
      showPage('laporan');
      if (!reportFrom.value && !reportTo.value) setReportPeriod('month');
      else {
        setReportMode(reportMonth.value ? 'month' : 'range');
        refreshReportDashboard();
      }
    }
    function setReportPeriod(period) {
      clearReportRankingFilter(false);
      const today = new Date();
      if (period === 'today') {
        setReportMode('range');
        reportFrom.value = localDateText(today);
        reportTo.value = localDateText(today);
        reportMonth.value = '';
      } else if (period === 'month') {
        reportMonth.value = `${today.getFullYear()}-${String(today.getMonth() + 1).padStart(2, '0')}`;
        setReportMode('month');
        reportFrom.value = localDateText(new Date(today.getFullYear(), today.getMonth(), 1));
        reportTo.value = localDateText(new Date(today.getFullYear(), today.getMonth() + 1, 0));
      } else {
        setReportMode('range');
        reportFrom.value = '';
        reportTo.value = '';
        reportMonth.value = '';
      }
      refreshReportDashboard();
    }
    function useCustomReportDates() {
      clearReportRankingFilter(false);
      reportMonth.value = '';
      setReportMode('range');
      refreshReportDashboard();
    }
    function setReportMonth() {
      if (!reportMonth.value) return;
      clearReportRankingFilter(false);
      setReportMode('month');
      const [year, month] = reportMonth.value.split('-').map(Number);
      reportFrom.value = `${year}-${String(month).padStart(2, '0')}-01`;
      reportTo.value = localDateText(new Date(year, month, 0));
      refreshReportDashboard();
    }
    const escapeHtml = (value) => String(value ?? '')
      .replaceAll('&', '&amp;').replaceAll('<', '&lt;').replaceAll('>', '&gt;')
      .replaceAll('"', '&quot;').replaceAll("'", '&#039;');
    function setReportStatus(message, isError=false, autoHide=false) {
      clearTimeout(reportStatusTimer);
      reportStatus.textContent = message;
      reportStatus.classList.toggle('bad', isError);
      if (autoHide) reportStatusTimer = setTimeout(() => { reportStatus.textContent = ''; }, 1600);
    }
    function renderRanking(id, rows, nameKey, showBrand=false, filterKind='') {
      const el = document.getElementById(id);
      if (!rows.length) {
        el.innerHTML = '<div class="empty-state">Belum ada penjualan pada periode ini.</div>';
        return;
      }
      el.innerHTML = rows.map((item, index) => {
        const value = String(item[nameKey] ?? '');
        const active = reportRankingFilter.kind === filterKind && reportRankingFilter.value === value;
        return `
        <button type="button" class="ranking-row ${active ? 'active' : ''}" onclick="setReportRankingFilter('${filterKind}','${encodeURIComponent(value)}')">
          <span class="ranking-number">${index + 1}</span>
          <span><span class="ranking-name">${escapeHtml(value)}</span>${showBrand ? `<span class="ranking-meta">${escapeHtml(item.brand)}</span>` : ''}</span>
          <span class="ranking-count">${item.count} unit</span>
        </button>`;
      }).join('');
    }
    function renderReportRankingViews() {
      renderRanking('reportTopBrands', reportTopBrandsRows, 'brand', false, 'brand');
      renderRanking('reportTopTypes', reportTopTypesRows, 'tipe', true, 'tipe');
    }
    function updateReportFilterChip() {
      if (!reportRankingFilter.value) {
        reportActiveFilter.classList.add('hidden');
        reportActiveFilter.innerHTML = '';
        return;
      }
      const label = reportRankingFilter.kind === 'brand' ? 'Merek' : 'Tipe';
      reportActiveFilter.innerHTML = `<span>${label}: ${escapeHtml(reportRankingFilter.value)}</span><button type="button" onclick="clearReportRankingFilter()" aria-label="Hapus filter">×</button>`;
      reportActiveFilter.classList.remove('hidden');
    }
    function setReportRankingFilter(kind, encodedValue) {
      const value = decodeURIComponent(encodedValue);
      if (reportRankingFilter.kind === kind && reportRankingFilter.value === value) reportRankingFilter = {kind:'', value:''};
      else reportRankingFilter = {kind, value};
      updateReportFilterChip();
      renderReportRankingViews();
      renderReportTransactions();
    }
    function clearReportRankingFilter(shouldRender=true) {
      reportRankingFilter = {kind:'', value:''};
      if (!shouldRender || !document.getElementById('reportActiveFilter')) return;
      updateReportFilterChip();
      renderReportRankingViews();
      renderReportTransactions();
    }
    function toggleCompactDetail(prefix, id, button) {
      const row = document.getElementById(prefix + id);
      const opening = row.classList.contains('hidden');
      row.classList.toggle('hidden', !opening);
      button.textContent = opening ? 'Tutup' : 'Detail';
      button.setAttribute('aria-expanded', opening ? 'true' : 'false');
    }
    function toggleReportDetail(id, button) {
      toggleCompactDetail('reportDetail-', id, button);
    }
    function renderReportTransactions() {
      const query = (reportSearch.value || '').trim().toLowerCase();
      const rows = reportSalesRows.filter(row => {
        const note = row.batch_id || ('NT-' + String(row.id).padStart(6, '0'));
        const matchesSearch = !query || [note, row.id, row.imei, row.tipe, row.brand, row.customer, row.catatan]
          .some(value => String(value ?? '').toLowerCase().includes(query));
        const matchesRanking = !reportRankingFilter.value
          || (reportRankingFilter.kind === 'brand' && row.brand === reportRankingFilter.value)
          || (reportRankingFilter.kind === 'tipe' && row.tipe === reportRankingFilter.value);
        return matchesSearch && matchesRanking;
      });
      reportResultCount.textContent = `${rows.length} transaksi`;
      const header = '<tr><th>Tanggal</th><th>Tipe barang</th><th class="cell-right">Harga jual</th><th>Customer</th><th>Detail</th></tr>';
      if (!rows.length) {
        const message = query || reportRankingFilter.value ? 'Tidak ada transaksi yang cocok.' : 'Belum ada transaksi pada periode ini.';
        reportSalesTable.innerHTML = header + `<tr><td colspan="5" class="empty-state">${message}</td></tr>`;
        return;
      }
      reportSalesTable.innerHTML = header + rows.map(row => {
        const note = row.batch_id || ('NT-' + String(row.id).padStart(6, '0'));
        return `
          <tr>
            <td class="cell-nowrap">${escapeHtml(formatReportDateTime(row.tanggal_jual))}</td>
            <td><b>${escapeHtml(row.tipe)}</b><div class="ranking-meta">${escapeHtml(row.brand)}</div></td>
            <td class="cell-right">${escapeHtml(rupiah(row.harga_jual))}</td>
            <td>${escapeHtml(row.customer || '-')}</td>
            <td><button type="button" class="report-detail-button" aria-expanded="false" onclick="toggleReportDetail(${Number(row.id)},this)">Detail</button></td>
          </tr>
          <tr id="reportDetail-${Number(row.id)}" class="report-detail-row hidden"><td colspan="5">
            <div class="report-detail-grid">
              <div class="report-detail-item"><div class="report-detail-label">Nomor nota</div><div class="report-detail-value">${escapeHtml(note)}</div></div>
              <div class="report-detail-item"><div class="report-detail-label">IMEI</div><div class="report-detail-value">${escapeHtml(row.imei)}</div></div>
              <div class="report-detail-item"><div class="report-detail-label">Merek</div><div class="report-detail-value">${escapeHtml(row.brand)}</div></div>
              <div class="report-detail-item"><div class="report-detail-label">HPP</div><div class="report-detail-value">${escapeHtml(rupiah(row.hpp))}</div></div>
              <div class="report-detail-item"><div class="report-detail-label">Laba/rugi</div><div class="report-detail-value">${escapeHtml(rupiah(row.laba_rugi))}</div></div>
              <div class="report-detail-item"><div class="report-detail-label">Catatan</div><div class="report-detail-value">${escapeHtml(row.catatan || '-')}</div></div>
            </div>
          </td></tr>`;
      }).join('');
    }
    async function refreshReportDashboard() {
      try {
        updateReportPeriodLabel();
        setReportStatus('Memuat laporan…');
        const params = new URLSearchParams({from:reportFrom.value || '', to:reportTo.value || ''});
        const [summaryData, salesData] = await Promise.all([
          api('/api/summary?' + params.toString()),
          api('/api/sales-history?' + params.toString()),
        ]);
        reportUnits.textContent = `${summaryData.jumlah_terjual_periode} unit`;
        reportSales.textContent = rupiah(summaryData.total_jual);
        reportTopBrandsRows = summaryData.top_brands || [];
        reportTopTypesRows = summaryData.top_types || [];
        reportSalesRows = salesData.sales || [];
        updateReportFilterChip();
        renderReportRankingViews();
        renderReportTransactions();
        setReportStatus('Laporan diperbarui', false, true);
      } catch(e) {
        reportSalesTable.innerHTML = `<tr><td class="bad">${e.message}</td></tr>`;
        reportResultCount.textContent = '0 transaksi';
        setReportStatus(e.message, true);
      }
    }
    async function refreshSummary() {
      if (document.getElementById('page-laporan').classList.contains('active')) await refreshReportDashboard();
    }
    async function loadTables() {
      return;
    }
    function setStockStatus(status) {
      stockStatus = status;
      btnBelum.classList.toggle('active', status === 'BELUM_LAKU');
      btnTerjual.classList.toggle('active', status === 'TERJUAL');
      selectedBrand = '';
      loadGroups();
    }
    async function loadGroups() {
      if (!document.getElementById('groupList')) return;
      const data = await api('/api/groups?status=' + encodeURIComponent(stockStatus));
      const groups = data.groups || [];
      groupList.innerHTML = groups.map((g, i) => `
        <button class="group-btn ${(!selectedBrand && i===0) || selectedBrand===g.brand ? 'active' : ''}" onclick="loadTypes('${g.brand.replaceAll("'", "\\'")}')">
          <span><b>${g.brand}</b><br><span class="muted">Klik untuk lihat tipe</span></span>
          <span class="pill">${g.count} pcs</span>
        </button>
      `).join('') || '<div class="muted">Belum ada data.</div>';
      if (groups.length && !selectedBrand) await loadTypes(groups[0].brand);
      if (!groups.length) { typeList.innerHTML=''; detailTitle.textContent='Detail jenis'; detailPill.textContent='0 pcs'; }
    }
    async function loadTypes(brand) {
      selectedBrand = brand;
      document.querySelectorAll('.group-btn').forEach(btn => btn.classList.toggle('active', btn.textContent.includes(brand)));
      const data = await api('/api/types?status=' + encodeURIComponent(stockStatus) + '&brand=' + encodeURIComponent(brand));
      const types = data.types || [];
      const total = types.reduce((sum, row) => sum + row.count, 0);
      detailTitle.textContent = 'Detail ' + brand;
      detailPill.textContent = total + ' pcs';
      typeList.innerHTML = types.map(t => `
        <button class="group-btn" onclick="loadImeis('${t.tipe.replaceAll("'", "\\'")}')">
          <span><b>${t.tipe}</b><br><span class="muted">Klik untuk lihat IMEI</span></span>
          <span class="pill">${t.count} pcs</span>
        </button>
      `).join('') || '<div class="muted">Tidak ada tipe.</div>';
    }
    async function loadImeis(tipe) {
      const data = await api('/api/type-items?status=' + encodeURIComponent(stockStatus) + '&tipe=' + encodeURIComponent(tipe));
      const rows = data.items || [];
      const sold = stockStatus === 'TERJUAL';
      modalTitle.textContent = tipe;
      modalSubtitle.textContent = `${sold ? 'Terjual' : 'Tersedia'} • ${rows.length} pcs`;
      if (sold) {
        modalBody.innerHTML = `<table><tr><th>IMEI</th><th>Grade</th><th>Spek</th><th>HPP</th><th>Harga jual</th><th>Laba/Rugi</th><th>Tanggal jual</th><th>Nota</th></tr>${
          rows.map(r => {
            const nota = r.batch_id || (r.sale_id ? 'NT-' + String(r.sale_id).padStart(6, '0') : '-');
            return `<tr><td>${r.imei}</td><td>${r.grade || '-'}</td><td>${r.spek || '-'}</td><td>${rupiah(r.hpp)}</td><td>${rupiah(r.harga_jual)}</td><td>${rupiah(r.laba_rugi)}</td><td>${r.tanggal_jual || '-'}</td><td>${nota}</td></tr>`;
          }).join('')
        }</table>`;
      } else {
        modalBody.innerHTML = `<table><tr><th>IMEI</th><th>Grade</th><th>Spek</th><th>HPP</th></tr>${
          rows.map(r => `<tr><td>${r.imei}</td><td>${r.grade || '-'}</td><td>${r.spek || '-'}</td><td>${rupiah(r.hpp)}</td></tr>`).join('')
        }</table>`;
      }
      imeiModal.classList.add('active');
    }
    function closeModal() {
      imeiModal.classList.remove('active');
    }
    async function uploadFile(kind) {
      const fileEl = fileDataMasuk;
      const resultEl = resultDataMasuk;
      const f = fileEl.files[0];
      if (!f) return resultEl.textContent = 'Pilih file dulu.';
      resultEl.textContent = 'Mengupload...';
      const buffer = await f.arrayBuffer();
      let binary = '';
      const bytes = new Uint8Array(buffer);
      for (let i = 0; i < bytes.length; i++) binary += String.fromCharCode(bytes[i]);
      const payload = { filename:f.name, data:btoa(binary), import_type:kind, supplier:inSupplier.value, invoice:inInvoice.value, tanggal_pembelian:inDate.value };
      try {
        const result = await api('/api/import', { method:'POST', headers:{'Content-Type':'application/json'}, body: JSON.stringify(payload) });
        resultEl.innerHTML = `<span class="ok">Berhasil.</span> Baris: ${result.jumlah_baris}, baru: ${result.barang_baru}, update: ${result.barang_update}`;
        await refreshSummary(); await loadGroups(); await loadTables(); await loadInboundHistory();
      } catch(e) { resultEl.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    async function cekImei() {
      const imei = scanImei.value.trim();
      if (!imei) return;
      const data = await api('/api/item?imei=' + encodeURIComponent(imei));
      if (!data.items || data.items.length === 0) {
        modalTitle.textContent = 'IMEI tidak ditemukan';
        modalSubtitle.textContent = imei;
        modalBody.innerHTML = '<p class="bad">Barang tidak ditemukan di database.</p>';
        imeiModal.classList.add('active');
        return;
      }
      if (data.items.length > 1) {
        modalTitle.textContent = 'Pilih IMEI yang benar';
        modalSubtitle.textContent = 'Input ' + imei + ' cocok ke lebih dari satu barang';
        modalBody.innerHTML = data.items.map(b => `<div class="item-choice"><b>${b.tipe}</b><br>IMEI: ${b.imei}<br>HPP: ${rupiah(b.hpp)}<br>Status: ${b.status}<button onclick="pilihImei('${b.imei}'); closeModal()">Pilih IMEI ini</button></div>`).join('');
        imeiModal.classList.add('active');
        return;
      }
      pilihDariBarang(data.items[0]);
    }
    function pilihImei(imei) {
      scanImei.value = imei;
      api('/api/item?imei=' + encodeURIComponent(imei)).then(data => pilihDariBarang(data.items[0]));
    }
    function pilihDariBarang(b) {
      jualImei.value = b.imei;
      showPage('jual');
      cekJualSatuan();
    }
    async function uploadPaste() {
      const text = pasteData.value.trim();
      if (!text) return pasteResult.textContent = 'Paste data dulu.';
      const result = await api('/api/import-text', { method:'POST', headers:{'Content-Type':'application/json'}, body: JSON.stringify({ text, supplier:inSupplier.value, invoice:inInvoice.value, tanggal_pembelian:inDate.value }) });
      pasteResult.innerHTML = `<span class="ok">Berhasil.</span> Baris: ${result.jumlah_baris}, baru: ${result.barang_baru}, update: ${result.barang_update}`;
      pasteData.value = '';
      await refreshSummary(); await loadGroups(); await loadInboundHistory();
    }
    function renderJualDetail(item) {
      jualSatuanItem = item;
      const price = parseNumberInput(hargaJual.value);
      const laba = price != null && item.hpp != null ? price - item.hpp : null;
      jualDetail.style.display = 'block';
      jualDetail.innerHTML = `
        <div><span class="pill">${item.status}</span></div>
        <p><b>${item.tipe}</b><br>IMEI: ${item.imei}<br>HPP internal: ${rupiah(item.hpp)}<br>Grade: ${item.grade || '-'} | Spek: ${item.spek || '-'}</p>
        <div class="muted">Laba/rugi internal: <b>${rupiah(laba)}</b></div>
      `;
    }
    async function cekJualSatuan() {
      const imei = jualImei.value.trim();
      jualSatuanItem = null;
      jualDetail.style.display = 'none';
      jualResult.textContent = '';
      if (!imei) return jualResult.innerHTML = '<span class="bad">Isi IMEI dulu.</span>';
      try {
        const data = await api('/api/item?imei=' + encodeURIComponent(imei));
        const items = (data.items || []).filter(item => item.status === 'BELUM_LAKU');
        if (!items.length) {
          jualResult.innerHTML = '<span class="bad">IMEI tidak ditemukan atau sudah terjual.</span>';
          return;
        }
        if (items.length > 1) {
          jualDetail.style.display = 'block';
          jualDetail.innerHTML = '<p class="bad">6 digit cocok ke lebih dari satu barang. Pilih yang benar:</p>' +
            items.map(item => `<div class="item-choice"><b>${item.tipe}</b><br>IMEI: ${item.imei}<br>HPP: ${rupiah(item.hpp)}<button onclick="jualImei.value='${item.imei}'; cekJualSatuan()">Pilih</button></div>`).join('');
          return;
        }
        jualImei.value = items[0].imei;
        renderJualDetail(items[0]);
      } catch (e) {
        jualResult.innerHTML = `<span class="bad">${e.message}</span>`;
      }
    }
    hargaJual.addEventListener('input', () => {
      if (jualSatuanItem) renderJualDetail(jualSatuanItem);
    });
    async function jualBarang() {
      try {
        if (!jualSatuanItem || jualSatuanItem.imei !== jualImei.value.trim()) {
          await cekJualSatuan();
        }
        if (!jualSatuanItem) throw new Error('Cek detail barang dulu sebelum submit.');
        const result = await api('/api/sell', { method:'POST', headers:{'Content-Type':'application/json'}, body: JSON.stringify({ imei:jualImei.value, harga_jual:hargaJual.value, customer:customer.value, catatan:catatan.value }) });
        jualResult.innerHTML = `<span class="ok">Transaksi tersimpan.</span> Nota sedang dibuka. Laba/rugi internal: ${rupiah(result.laba_rugi)}<br><button onclick="window.open('/nota?id=${result.id}', '_blank')">Buka Nota Lagi</button>`;
        window.open('/nota?id=' + result.id, '_blank');
        hargaJual.value = ''; customer.value = ''; catatan.value = ''; jualImei.value = '';
        jualSatuanItem = null;
        jualDetail.style.display = 'none';
        await refreshSummary(); await loadGroups(); await loadTables();
      } catch (e) { jualResult.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    function parseNumberInput(value) {
      if (value == null) return null;
      const cleaned = String(value).replace(/[^\d,.-]/g, '').replace(/\./g, '').replace(',', '.');
      if (!cleaned) return null;
      const n = Number(cleaned);
      return Number.isFinite(n) ? n : null;
    }
    function updateBulkTotals() {
      let canSubmit = bulkReviewItems.length > 0;
      let totalJual = 0;
      let totalLaba = 0;
      bulkReviewItems.forEach((item, idx) => {
        const input = document.getElementById('bulkPrice' + idx);
        const price = parseNumberInput(input?.value);
        item.harga_jual = price;
        const priceRequiredStatus = item.status === 'OK' || item.status === 'BUTUH_HARGA';
        if (item.status === 'BUTUH_HARGA' && price != null) item.status = 'OK';
        if (!priceRequiredStatus || price == null) canSubmit = false;
        const laba = price != null && item.hpp != null ? price - item.hpp : null;
        item.laba_rugi = laba;
        const labaEl = document.getElementById('bulkLaba' + idx);
        if (labaEl) labaEl.textContent = rupiah(laba);
        const statusEl = document.getElementById('bulkStatus' + idx);
        if (statusEl && priceRequiredStatus) {
          statusEl.innerHTML = price == null ? '<span class="bad">Isi harga jual</span>' : '<span class="ok">OK</span>';
        }
        if (price != null) totalJual += price;
        if (laba != null) totalLaba += laba;
      });
      bulkResult.innerHTML = bulkReviewItems.length
        ? `Total jual: <b>${rupiah(totalJual)}</b> • Laba internal: <b>${rupiah(totalLaba)}</b>`
        : '';
      bulkSubmitBtn.disabled = !canSubmit;
    }
    function renderBulkReview(items) {
      bulkReviewItems = items || [];
      if (!bulkReviewItems.length) {
        bulkReview.innerHTML = '<p class="muted">Belum ada barang untuk direview.</p>';
        bulkSubmitBtn.disabled = true;
        return;
      }
      bulkReview.innerHTML = `<table>
        <tr><th>Input</th><th>IMEI</th><th>Tipe</th><th>HPP</th><th>Harga jual</th><th>Laba/Rugi</th><th>Status</th></tr>
        ${bulkReviewItems.map((item, idx) => `
          <tr>
            <td>${item.input || ''}</td>
            <td>${item.imei || '-'}</td>
            <td>${item.tipe || '-'}</td>
            <td>${rupiah(item.hpp)}</td>
            <td><input id="bulkPrice${idx}" value="${item.harga_jual ?? ''}" placeholder="Harga jual" oninput="updateBulkTotals()"></td>
            <td id="bulkLaba${idx}">${rupiah(item.laba_rugi)}</td>
            <td id="bulkStatus${idx}">${item.status === 'OK' ? '<span class="ok">OK</span>' : '<span class="bad">' + (item.message || item.status) + '</span>'}</td>
          </tr>
          ${item.choices && item.choices.length ? `<tr><td colspan="7"><span class="bad">Pilihan IMEI:</span> ${item.choices.map(c => `${c.imei} (${c.tipe})`).join(', ')}</td></tr>` : ''}
        `).join('')}
      </table>`;
      updateBulkTotals();
    }
    async function reviewMassal() {
      try {
        bulkResult.textContent = 'Mengecek barang...';
        const result = await api('/api/review-bulk-sale', {
          method:'POST',
          headers:{'Content-Type':'application/json'},
          body: JSON.stringify({ items_text: bulkImeis.value, harga_jual_sama: bulkHargaSama.value })
        });
        renderBulkReview(result.items);
      } catch (e) {
        bulkSubmitBtn.disabled = true;
        bulkResult.innerHTML = `<span class="bad">${e.message}</span>`;
      }
    }
    async function jualMassal() {
      try {
        bulkResult.innerHTML = 'Menyimpan transaksi...';
        updateBulkTotals();
        if (bulkSubmitBtn.disabled) throw new Error('Review belum OK atau harga jual masih kosong.');
        const result = await api('/api/sell-bulk', {
          method:'POST',
          headers:{'Content-Type':'application/json'},
          body: JSON.stringify({
            customer: bulkCustomer.value,
            items: bulkReviewItems.map(item => ({ imei: item.imei, harga_jual: item.harga_jual })),
            catatan: bulkCatatan.value
          })
        });
        bulkResult.innerHTML = `<span class="ok">Penjualan massal tersimpan.</span> Nota batch sedang dibuka. Jumlah: ${result.jumlah} pcs, total: ${rupiah(result.total_jual)}<br><button onclick="window.open('/nota-batch?id=${result.batch_id}', '_blank')">Buka Nota Batch Lagi</button>`;
        window.open('/nota-batch?id=' + result.batch_id, '_blank');
        bulkImeis.value = ''; bulkHargaSama.value = ''; bulkCustomer.value = ''; bulkCatatan.value = '';
        bulkReviewItems = [];
        bulkReview.innerHTML = '';
        bulkSubmitBtn.disabled = true;
        await refreshSummary(); await loadGroups(); await loadTables();
      } catch (e) { bulkResult.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    async function adminSearch() {
      try {
        adminResult.innerHTML = '<span class="muted">Mencari data...</span>';
        const data = await api('/api/admin/items?imei=' + encodeURIComponent(adminImei.value));
        const rows = data.items || [];
        adminResult.innerHTML = rows.map((r, idx) => {
          if (r.status === 'TERJUAL') {
            const nota = r.batch_id || (r.sale_id ? 'NT-' + String(r.sale_id).padStart(6, '0') : '-');
            return `<div class="upload-card" style="margin-top:10px">
              <b>${r.tipe} — ${r.imei}</b>
              <div class="muted">Status: Terjual • HPP ${rupiah(r.hpp)} • Nota ${nota} • ${r.tanggal_jual || '-'}</div>
              <div class="row"><input id="adminSalePrice${idx}" value="${r.harga_jual ?? ''}" placeholder="Harga jual"><input id="adminSaleReason${idx}" placeholder="Alasan koreksi / pembatalan"></div>
              <div class="row"><button onclick="adminCorrectSale(${r.sale_id}, ${idx})">Koreksi harga jual</button><button class="danger" onclick="adminVoidSale(${r.sale_id}, ${idx})">Batalkan transaksi</button></div>
            </div>`;
          }
          return `<div class="upload-card" style="margin-top:10px">
            <b>Barang belum terjual — ${r.imei}</b>
            <div class="row"><input id="adminEditImei${idx}" value="${r.imei}" placeholder="IMEI"><input id="adminEditTipe${idx}" value="${r.tipe || ''}" placeholder="Tipe"></div>
            <div class="row"><input id="adminEditHpp${idx}" value="${r.hpp ?? ''}" placeholder="HPP"><input id="adminEditModal${idx}" value="${r.modal_awal ?? ''}" placeholder="Modal awal"></div>
            <div class="row"><input id="adminEditGrade${idx}" value="${r.grade || ''}" placeholder="Grade"><input id="adminEditDate${idx}" value="${(r.tanggal_masuk || '').slice(0,10)}" type="date"></div>
            <div class="row"><input id="adminEditSpek${idx}" value="${r.spek || ''}" placeholder="Spek"><input id="adminEditKet${idx}" value="${r.keterangan || ''}" placeholder="Keterangan"></div>
            <div class="row"><input id="adminEditSupplier${idx}" value="${r.supplier || ''}" placeholder="Supplier"><input id="adminEditInvoice${idx}" value="${r.invoice || ''}" placeholder="Invoice"></div>
            <div class="row"><button onclick="adminUpdateItem('${r.imei}', ${idx})">Simpan perubahan</button><button class="danger" onclick="adminDeleteItem('${r.imei}')">Hapus barang</button></div>
          </div>`;
        }).join('') || '<span class="bad">Data tidak ditemukan.</span>';
      } catch (e) { adminResult.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    async function adminUpdateItem(originalImei, idx) {
      try {
        if (!confirm('Simpan perubahan data barang ini?')) return;
        const result = await api('/api/admin/update-item', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({
          original_imei: originalImei, imei:document.getElementById('adminEditImei'+idx).value,
          tipe:document.getElementById('adminEditTipe'+idx).value, hpp:document.getElementById('adminEditHpp'+idx).value,
          modal_awal:document.getElementById('adminEditModal'+idx).value, grade:document.getElementById('adminEditGrade'+idx).value,
          spek:document.getElementById('adminEditSpek'+idx).value, keterangan:document.getElementById('adminEditKet'+idx).value,
          tanggal_masuk:document.getElementById('adminEditDate'+idx).value, supplier:document.getElementById('adminEditSupplier'+idx).value,
          invoice:document.getElementById('adminEditInvoice'+idx).value
        })});
        adminImei.value = result.imei; await adminSearch(); await refreshSummary(); await loadGroups();
      } catch(e) { alert(e.message); }
    }
    async function adminDeleteItem(imei) {
      try {
        if (!confirm('Hapus barang ' + imei + '? Backup akan dibuat otomatis.')) return;
        const result = await api('/api/admin/delete-item', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({imei})});
        adminResult.innerHTML = `<span class="ok">${result.message}</span><br><span class="muted">Backup: ${result.backup}</span>`;
        await refreshSummary(); await loadGroups();
      } catch(e) { alert(e.message); }
    }
    async function adminCorrectSale(saleId, idx) {
      try {
        const harga_jual = document.getElementById('adminSalePrice'+idx).value;
        const alasan = document.getElementById('adminSaleReason'+idx).value;
        if (!confirm('Koreksi harga jual transaksi ini?')) return;
        const result = await api('/api/admin/update-sale', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({sale_id:saleId,harga_jual,alasan})});
        await adminSearch(); await refreshSummary(); await loadTables();
      } catch(e) { alert(e.message); }
    }
    async function adminVoidSale(saleId, idx) {
      try {
        const alasan = document.getElementById('adminSaleReason'+idx).value;
        if (!confirm('Batalkan transaksi? Barang akan kembali menjadi stok.')) return;
        const result = await api('/api/admin/void-sale', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({sale_id:saleId,alasan})});
        adminResult.innerHTML = `<span class="ok">${result.message}</span><br><span class="muted">Backup: ${result.backup}</span>`;
        await refreshSummary(); await loadGroups(); await loadTables();
      } catch(e) { alert(e.message); }
    }
    async function resetDatabase() {
      try {
        if (resetConfirmation.value !== 'HAPUS SEMUA DATA') throw new Error('Ketik tepat HAPUS SEMUA DATA terlebih dahulu.');
        if (!confirm('PERINGATAN: seluruh stok dan transaksi akan dikosongkan. Lanjutkan?')) return;
        const result = await api('/api/admin/reset-database', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({konfirmasi:resetConfirmation.value})});
        resetResult.innerHTML = `<span class="ok">${result.message}</span><br>Backup tersimpan: ${result.backup}`;
        resetConfirmation.value = ''; adminResult.innerHTML = '';
        await refreshSummary(); selectedBrand=''; await loadGroups(); await loadTables();
      } catch(e) { resetResult.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    async function saveManualInbound() {
      try {
        const result = await api('/api/inbound-manual', {method:'POST', headers:{'Content-Type':'application/json'}, body:JSON.stringify({
          imei:manualImei.value, tipe:manualTipe.value, hpp:manualHpp.value, grade:manualGrade.value,
          spek:manualSpek.value, keterangan:manualKet.value, supplier:inSupplier.value,
          invoice:inInvoice.value, tanggal_pembelian:inDate.value
        })});
        manualInboundResult.innerHTML = `<span class="ok">Barang tersimpan.</span> Batch: ${result.batch_id}`;
        manualImei.value=''; manualTipe.value=''; manualHpp.value=''; manualGrade.value=''; manualSpek.value=''; manualKet.value='';
        await loadInboundHistory(); await loadInventory();
      } catch(e) { manualInboundResult.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    async function loadInboundHistory() {
      try {
        const data = await api('/api/inbound-history');
        renderTable('inboundTable', data.batches || [], [
          {key:'batch_id',label:'Batch'}, {key:'tanggal_pembelian',label:'Tanggal'}, {key:'supplier',label:'Supplier'},
          {key:'invoice',label:'Invoice'}, {key:'jumlah_baris',label:'Jumlah'}, {key:'total_modal',label:'Total modal',format:rupiah}
        ]);
      } catch(e) { inboundTable.innerHTML = `<tr><td class="bad">${e.message}</td></tr>`; }
    }
    async function loadSalesHistory() {
      try {
        const params = new URLSearchParams({q:salesQuery.value, from:salesFrom.value, to:salesTo.value});
        const data = await api('/api/sales-history?' + params.toString());
        renderSalesHistoryCompact(data.sales || []);
      } catch(e) {
        salesHistoryCount.textContent = '0 transaksi';
        salesHistoryTable.innerHTML = `<tr><td class="bad">${escapeHtml(e.message)}</td></tr>`;
      }
    }
    function scheduleSalesHistory() {
      clearTimeout(salesHistoryTimer);
      salesHistoryTimer = setTimeout(loadSalesHistory, 250);
    }
    function renderSalesHistoryCompact(rows) {
      salesHistoryCount.textContent = `${rows.length} transaksi`;
      const header = '<tr><th>Tanggal</th><th>Tipe barang</th><th class="cell-right">Harga jual</th><th>Customer</th><th>Detail</th></tr>';
      if (!rows.length) {
        salesHistoryTable.innerHTML = header + '<tr><td colspan="5" class="empty-state">Tidak ada transaksi yang cocok.</td></tr>';
        return;
      }
      salesHistoryTable.innerHTML = header + rows.map(row => {
        const id = Number(row.id);
        const note = row.batch_id || ('NT-' + String(id).padStart(6, '0'));
        const notaUrl = row.batch_id ? '/nota-batch?id=' + encodeURIComponent(row.batch_id) : '/nota?id=' + id;
        return `<tr>
          <td class="cell-nowrap">${escapeHtml(formatReportDateTime(row.tanggal_jual))}</td>
          <td><b>${escapeHtml(row.tipe || '-')}</b><div class="ranking-meta">${escapeHtml(row.brand || '-')}</div></td>
          <td class="cell-right">${escapeHtml(rupiah(row.harga_jual))}</td>
          <td>${escapeHtml(row.customer || '-')}</td>
          <td><button type="button" class="report-detail-button" aria-expanded="false" onclick="toggleCompactDetail('salesHistoryDetail-',${id},this)">Detail</button></td>
        </tr>
        <tr id="salesHistoryDetail-${id}" class="report-detail-row hidden"><td colspan="5">
          <div class="report-detail-grid">
            <div class="report-detail-item"><div class="report-detail-label">Nomor nota</div><div class="report-detail-value">${escapeHtml(note)}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">IMEI</div><div class="report-detail-value">${escapeHtml(row.imei)}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">HPP</div><div class="report-detail-value">${escapeHtml(rupiah(row.hpp))}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">Laba/rugi</div><div class="report-detail-value">${escapeHtml(rupiah(row.laba_rugi))}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">Catatan</div><div class="report-detail-value">${escapeHtml(row.catatan || '-')}</div></div>
            <div class="report-detail-item"><button class="ghost" onclick="window.open('${notaUrl}','_blank')">Cetak nota</button></div>
          </div>
        </td></tr>`;
      }).join('');
    }
    async function loadInventory() {
      try {
        const params = new URLSearchParams({q:inventoryQuery.value, status:inventoryStatus.value, brand:inventoryBrand.value});
        const data = await api('/api/inventory?' + params.toString());
        inventoryBrand.innerHTML = '<option value="">Semua merek</option>' + (data.brands || []).map(item =>
          `<option value="${item.brand}">${item.brand} (${item.count})</option>`
        ).join('');
        inventoryBrand.value = data.selected_brand || '';
        renderInventoryCompact(data.items || []);
      } catch(e) {
        inventoryCount.textContent = '0 barang';
        inventoryTable.innerHTML = `<tr><td class="bad">${escapeHtml(e.message)}</td></tr>`;
      }
    }
    function friendlyStatus(status) {
      return status === 'TERJUAL' ? 'Terjual' : 'Tersedia';
    }
    function renderInventoryCompact(rows) {
      inventoryCount.textContent = `${rows.length} barang`;
      const header = '<tr><th>IMEI</th><th>Tipe barang</th><th class="cell-right">HPP</th><th>Status</th><th>Detail</th></tr>';
      if (!rows.length) {
        inventoryTable.innerHTML = header + '<tr><td colspan="5" class="empty-state">Tidak ada barang yang cocok.</td></tr>';
        return;
      }
      inventoryTable.innerHTML = header + rows.map((row, index) => {
        const sold = row.status === 'TERJUAL';
        return `<tr>
          <td class="cell-nowrap">${escapeHtml(row.imei)}</td>
          <td><b>${escapeHtml(row.tipe || '-')}</b><div class="ranking-meta">${escapeHtml(row.brand || '-')}</div></td>
          <td class="cell-right">${escapeHtml(rupiah(row.hpp))}</td>
          <td><span class="status-badge ${sold ? 'status-sold' : 'status-available'}">${friendlyStatus(row.status)}</span></td>
          <td><button type="button" class="report-detail-button" aria-expanded="false" onclick="toggleCompactDetail('inventoryDetail-',${index},this)">Detail</button></td>
        </tr>
        <tr id="inventoryDetail-${index}" class="report-detail-row hidden"><td colspan="5">
          <div class="report-detail-grid">
            <div class="report-detail-item"><div class="report-detail-label">Grade</div><div class="report-detail-value">${escapeHtml(row.grade || '-')}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">Spek</div><div class="report-detail-value">${escapeHtml(row.spek || '-')}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">Tanggal masuk</div><div class="report-detail-value">${escapeHtml(formatReportDate(String(row.tanggal_masuk || '').slice(0,10)) || '-')}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">Umur stok</div><div class="report-detail-value">${row.umur_hari == null ? '-' : escapeHtml(row.umur_hari + ' hari')}</div></div>
            <div class="report-detail-item"><div class="report-detail-label">Keterangan</div><div class="report-detail-value">${escapeHtml(row.keterangan || '-')}</div></div>
          </div>
        </td></tr>`;
      }).join('');
    }
    async function loadImeiHistory() {
      try {
        const data = await api('/api/imei-history?imei=' + encodeURIComponent(historyImei.value));
        const b = data.item;
        imeiHistoryResult.innerHTML = `<div class="upload-card"><b>${escapeHtml(b.tipe)} — ${escapeHtml(b.imei)}</b><div class="muted">Status: ${friendlyStatus(b.status)} • HPP: ${escapeHtml(rupiah(b.hpp))}</div></div>` +
          `<div class="table-wrap"><table><tr><th>Tanggal</th><th>Kegiatan</th><th>Detail</th></tr>${(data.events||[]).map(e=>`<tr><td>${e.tanggal||''}</td><td>${e.jenis}</td><td>${e.detail||''}</td></tr>`).join('')}</table></div>`;
      } catch(e) { imeiHistoryResult.innerHTML = `<span class="bad">${e.message}</span>`; }
    }
    function downloadReport() {
      const params = new URLSearchParams({from:reportFrom.value || '', to:reportTo.value || ''});
      location.href = '/api/export?' + params.toString();
    }
    function renderTable(id, rows, cols, emptyText='') {
      const el = document.getElementById(id);
      const header = '<tr>' + cols.map(c => `<th class="${c.className || ''}">${c.label}</th>`).join('') + '</tr>';
      const body = rows.length
        ? rows.map(r => '<tr>' + cols.map(c => `<td class="${c.className || ''}">${c.format ? c.format(r[c.key], r) : (r[c.key] ?? '')}</td>`).join('') + '</tr>').join('')
        : (emptyText ? `<tr><td colspan="${cols.length}" class="empty-state">${emptyText}</td></tr>` : '');
      el.innerHTML = header + body;
    }
    loadInboundHistory();
  </script>
</body>
</html>
"""


class Handler(BaseHTTPRequestHandler):
    def send_json(self, payload: dict[str, Any], status: int = 200) -> None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self) -> dict[str, Any]:
        length = int(self.headers.get("Content-Length", "0"))
        raw = self.rfile.read(length)
        return json.loads(raw.decode("utf-8")) if raw else {}

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        try:
            if parsed.path == "/api/health":
                with db() as con:
                    con.execute("SELECT 1 AS ok").fetchone()
                self.send_json({"status": "ok", "database": "neon" if USE_POSTGRES else "sqlite"})
            elif parsed.path == "/":
                body = HTML.encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            elif parsed.path == "/assets/find-cell-logo.png":
                logo_path = APP_DIR / "assets" / "find-cell-logo.png"
                if not logo_path.exists():
                    self.send_json({"error": "Logo belum tersedia."}, 404)
                    return
                body = logo_path.read_bytes()
                self.send_response(200)
                self.send_header("Content-Type", "image/png")
                self.send_header("Cache-Control", "public, max-age=3600")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            elif parsed.path == "/api/summary":
                qs = parse_qs(parsed.query)
                self.send_json(summary(qs.get("from", [""])[0], qs.get("to", [""])[0]))
            elif parsed.path == "/api/item":
                qs = parse_qs(parsed.query)
                imei = clean_imei(qs.get("imei", [""])[0])
                self.send_json({"items": find_items_by_imei_input(imei, only_unsold=False, limit=20)})
            elif parsed.path == "/api/tables":
                self.send_json({"stok": list_barang("BELUM_LAKU", 200), "penjualan": list_penjualan(200)})
            elif parsed.path == "/api/inbound-history":
                self.send_json({"batches": list_purchase_batches(500)})
            elif parsed.path == "/api/sales-history":
                qs = parse_qs(parsed.query)
                self.send_json({"sales": list_penjualan(1000, qs.get("from", [""])[0], qs.get("to", [""])[0], qs.get("q", [""])[0])})
            elif parsed.path == "/api/inventory":
                qs = parse_qs(parsed.query)
                self.send_json(inventory_data(
                    qs.get("status", [""])[0],
                    qs.get("q", [""])[0],
                    qs.get("brand", [""])[0],
                    2000,
                ))
            elif parsed.path == "/api/imei-history":
                qs = parse_qs(parsed.query)
                self.send_json(imei_history(qs.get("imei", [""])[0]))
            elif parsed.path == "/api/groups":
                qs = parse_qs(parsed.query)
                status = qs.get("status", ["BELUM_LAKU"])[0]
                self.send_json({"groups": group_summary(status)})
            elif parsed.path == "/api/types":
                qs = parse_qs(parsed.query)
                status = qs.get("status", ["BELUM_LAKU"])[0]
                brand = qs.get("brand", ["iPhone"])[0]
                self.send_json({"types": type_summary(status, brand)})
            elif parsed.path == "/api/type-items":
                qs = parse_qs(parsed.query)
                status = qs.get("status", ["BELUM_LAKU"])[0]
                tipe = qs.get("tipe", [""])[0]
                self.send_json({"items": type_items(status, tipe)})
            elif parsed.path == "/api/admin/items":
                qs = parse_qs(parsed.query)
                imei = clean_imei(qs.get("imei", [""])[0])
                self.send_json({"items": admin_find_items(imei)})
            elif parsed.path == "/api/export":
                qs = parse_qs(parsed.query)
                body = export_excel(qs.get("from", [""])[0], qs.get("to", [""])[0])
                filename = f"resume-stok-imei-{datetime.now().strftime('%Y%m%d-%H%M%S')}.xlsx"
                self.send_response(200)
                self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
                self.send_header("Content-Disposition", f'attachment; filename="{filename}"')
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            elif parsed.path in {"/template/xlsx", "/template/csv"}:
                template_name = "template-upload-barang-masuk.xlsx" if parsed.path.endswith("xlsx") else "template-upload-barang-masuk.csv"
                template_path = APP_DIR / "templates" / template_name
                if not template_path.exists():
                    self.send_json({"error": "Template belum tersedia."}, 404)
                    return
                body = template_path.read_bytes()
                content_type = (
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    if template_name.endswith(".xlsx")
                    else "text/csv; charset=utf-8"
                )
                self.send_response(200)
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Disposition", f'attachment; filename="{template_name}"')
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            elif parsed.path == "/nota":
                qs = parse_qs(parsed.query)
                sale_id = int(qs.get("id", ["0"])[0] or "0")
                body = nota_html(sale_id).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            elif parsed.path == "/nota-batch":
                qs = parse_qs(parsed.query)
                batch_id = clean_text(qs.get("id", [""])[0])
                body = nota_batch_html(batch_id).encode("utf-8")
                self.send_response(200)
                self.send_header("Content-Type", "text/html; charset=utf-8")
                self.send_header("Content-Length", str(len(body)))
                self.end_headers()
                self.wfile.write(body)
            else:
                self.send_json({"error": "Halaman tidak ditemukan."}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 500)

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        try:
            if parsed.path == "/api/import":
                payload = self.read_json()
                filename = clean_text(payload.get("filename")) or "barang-masuk.xlsx"
                data = base64.b64decode(payload.get("data", ""))
                suffix = Path(filename).suffix or ".xlsx"
                with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
                    tmp.write(data)
                    temp_path = Path(tmp.name)
                try:
                    records = parse_upload_file(temp_path, filename)
                    result = import_records(records, filename, payload)
                finally:
                    try:
                        temp_path.unlink()
                    except OSError:
                        pass
                self.send_json(result)
            elif parsed.path == "/api/import-text":
                payload = self.read_json()
                records = parse_text_rows(payload.get("text", ""), "Paste manual")
                payload["import_type"] = "DATA_MASUK"
                result = import_records(records, "Paste manual", payload)
                self.send_json(result)
            elif parsed.path == "/api/inbound-manual":
                self.send_json(manual_inbound(self.read_json()))
            elif parsed.path == "/api/sell":
                result = sell_item(self.read_json())
                self.send_json(result)
            elif parsed.path == "/api/review-bulk-sale":
                result = review_bulk_sale(self.read_json())
                self.send_json(result)
            elif parsed.path == "/api/sell-bulk":
                result = sell_bulk(self.read_json())
                self.send_json(result)
            elif parsed.path == "/api/admin/update-item":
                self.send_json(admin_update_item(self.read_json()))
            elif parsed.path == "/api/admin/delete-item":
                self.send_json(admin_delete_item(self.read_json()))
            elif parsed.path == "/api/admin/update-sale":
                self.send_json(admin_update_sale(self.read_json()))
            elif parsed.path == "/api/admin/void-sale":
                self.send_json(admin_void_sale(self.read_json()))
            elif parsed.path == "/api/admin/reset-database":
                self.send_json(admin_reset_database(self.read_json()))
            else:
                self.send_json({"error": "Endpoint tidak ditemukan."}, 404)
        except Exception as exc:
            self.send_json({"error": str(exc)}, 400)

    def log_message(self, format: str, *args: Any) -> None:
        print(f"[{now_text()}] {self.address_string()} {format % args}")


def main() -> None:
    init_db()
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Sistem berjalan di http://{HOST}:{PORT}")
    print("Tekan Ctrl+C untuk berhenti.")
    server.serve_forever()


if USE_POSTGRES:
    init_db()

handler = Handler


if __name__ == "__main__":
    main()
